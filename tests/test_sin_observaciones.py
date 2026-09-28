from io import BytesIO
import unittest
from docx import Document
from test_asistencia_tecnica_word import generar, solicitud


class SinObservacionesTest(unittest.TestCase):
    def bloque(self, fila):
        doc = Document(BytesIO(generar([fila])))
        textos = [p.text.strip() for p in doc.paragraphs]
        inicio = textos.index('Observaciones respecto de se ajusta Si/No:') + 1
        fin = next(i for i in range(inicio, len(textos)) if textos[i].startswith('En cuanto a la información'))
        return doc, [t for t in textos[inicio:fin] if t]

    def test_validacion_completa_sin_diferencias(self):
        doc, bloque = self.bloque(solicitud())
        self.assertEqual(bloque, ['• Sin observaciones'])
        p = next(p for p in doc.paragraphs if 'Sin observaciones' in p.text)
        self.assertTrue(all(r.font.name == 'Garamond' for r in p.runs if r.text))
        self.assertTrue(all(r.font.size.pt == 11 for r in p.runs if r.text))

    def test_diferencia_conserva_observacion(self):
        _, bloque = self.bloque(solicitud(valor_adicion=5000000))
        self.assertTrue(any('Adición: Sería por' in t for t in bloque))
        self.assertFalse(any('Sin observaciones' in t for t in bloque))

    def test_datos_ausentes_no_se_declaran_correctos(self):
        for campo in ('contrato', 'contratista', 'objeto', 'fecha_inicio', 'plazo_inicial',
                      'fecha_terminacion_inicial', 'fecha_terminacion_final', 'prorroga_solicitada',
                      'valor_inicial', 'valor_adicion'):
            with self.subTest(campo=campo):
                _, bloque = self.bloque(solicitud(**{campo: ''}))
                self.assertFalse(any('Sin observaciones' in t for t in bloque))
                self.assertTrue(bloque)

    def test_datos_invalidos_no_se_declaran_correctos(self):
        _, bloque = self.bloque(solicitud(fecha_inicio='fecha desconocida'))
        self.assertEqual(bloque, ['• Pendiente de revisión: faltan datos para completar la validación.'])

    def test_observacion_vigencia_impide_sin_observaciones(self):
        _, bloque = self.bloque(solicitud(prorroga_solicitada='4 meses', fecha_terminacion_final='30/01/2027', valor_adicion=12000000))
        self.assertTrue(any('vigencia fiscal' in t for t in bloque))
        self.assertFalse(any('Sin observaciones' in t for t in bloque))
