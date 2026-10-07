from io import BytesIO
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook
from asistencia_tecnica import (
    analizar_solicitudes, analizar_solicitud_pdf, _limpiar_prorroga_solicitada,
    _extraer_prorroga_numeral_dos, _datos_calculadora, generar_excel_asistencia,
    calcular_fecha_fin_prorroga,
    _nombre_word_localidad, _nombre_excel_asistencia, generar_zip_asistencia,
)
from test_asistencia_tecnica_word import solicitud, generar, PLANTILLA, FECHA_RESPUESTA


class ProrrogaSinUnidadTest(unittest.TestCase):
    def test_cantidad_sola_es_dias_sin_convertir_fechas_o_dinero(self):
        for original, esperado in [("84", "84 días"), (" 084 ", "84 días"),
                                  (1, "1 día"), ("0", "0 días"),
                                  ("8 meses", "8 meses"), ("31/12/2026", "31/12/2026"),
                                  ("$84", "$84"), ("84.000", "84.000"), ("N/A", "")]:
            self.assertEqual(_limpiar_prorroga_solicitada(original), esperado)

    def test_extrae_solo_el_campo_tiempo_del_numeral_dos(self):
        for texto in ("84", "\n84\n", "84 días"):
            fuente = "II. INFORMACIÓN DE LA MODIFICACIÓN SOLICITADA Tiempo: " + texto + " Fecha Terminación Final: 31/12/2026"
            self.assertEqual(_extraer_prorroga_numeral_dos(fuente), "84 días")
        self.assertEqual(_extraer_prorroga_numeral_dos(
            "II. INFORMACIÓN DE LA MODIFICACIÓN SOLICITADA Prórroga Tiempo 84 Fecha Terminación Final: 31/12/2026"), "84 días")
        self.assertEqual(_extraer_prorroga_numeral_dos(
            "II. INFORMACIÓN DE LA MODIFICACIÓN SOLICITADA Adición Valor: $84 "
            "III. INFORMACIÓN DE MODIFICACIONES ANTERIORES Tiempo: 84"), "")

    def test_word_excel_calculadora_incluyen_dias_en_datos_editados(self):
        fila = solicitud(prorroga_solicitada="84")
        doc = Document(BytesIO(generar([fila])))
        self.assertEqual(doc.tables[0].rows[9].cells[1].text, "84 días")
        datos = _datos_calculadora(fila)
        self.assertEqual((datos["prorroga_meses"], datos["prorroga_dias"]), (0, 84))
        self.assertEqual(calcular_fecha_fin_prorroga(date(2026, 8, 1), "84"), date(2026, 10, 24))
        wb = load_workbook(BytesIO(generar_excel_asistencia([fila])))
        self.assertTrue(any(c.value == "84 días" for row in wb["CPS 2026"] for c in row))

    def test_analisis_pdf_transmite_cantidad_a_la_fila(self):
        texto = "I. RESUMEN CONTRACTUAL Objeto: Prueba II. INFORMACIÓN DE LA MODIFICACIÓN SOLICITADA Tiempo: 84 Fecha Terminación Final: 31/12/2026"
        with patch("asistencia_tecnica.extraer_texto_pdf", return_value=(texto, [])):
            fila, _ = analizar_solicitud_pdf("164717.pdf", b"pdf")
        self.assertEqual(fila["prorroga_solicitada"], "84 días")


class LoteSolicitudesTest(unittest.TestCase):
    def test_lote_de_35_sin_limite_y_fallo_aislado(self):
        def leer(nombre, contenido):
            if nombre == "15.pdf":
                raise ValueError("PDF defectuoso")
            return {"archivo": nombre}, []
        progreso = []
        with patch("asistencia_tecnica.analizar_solicitud_pdf", side_effect=leer) as analizar:
            filas, errores = analizar_solicitudes(
                ((f"{n}.pdf", b"PDF") for n in range(35)),
                progreso=lambda n, nombre: progreso.append(n),
            )
        self.assertEqual(analizar.call_count, 35)
        self.assertEqual(len(filas), 35)
        self.assertEqual(filas[-1]["archivo"], "34.pdf")
        self.assertIn("15.pdf", errores[0])
        self.assertEqual(progreso[-1], 35)

    def test_interrupcion_retoma_sin_leer_otra_vez_y_detecta_pdf_cambiado(self):
        cache = {}
        def interrumpir(n, nombre):
            if n == 1:
                raise KeyboardInterrupt()
        archivos = [("uno.pdf", b"uno"), ("dos.pdf", b"dos")]
        with patch("asistencia_tecnica.analizar_solicitud_pdf", return_value=({"dato": "original"}, [])) as leer:
            with self.assertRaises(KeyboardInterrupt):
                analizar_solicitudes(iter(archivos), cache, interrumpir)
            filas, _ = analizar_solicitudes(iter(archivos), cache)
            self.assertEqual(leer.call_count, 2)
            filas[0]["dato"] = "editado"
            self.assertEqual(analizar_solicitudes(iter(archivos), cache)[0][0]["dato"], "original")
            analizar_solicitudes([("uno.pdf", b"cambio")], cache)
            self.assertEqual(leer.call_count, 3)

    def test_lectura_perezosa_no_prepara_todos_los_pdf(self):
        eventos = []
        def archivos():
            for n in range(3):
                eventos.append(f"carga{n}")
                yield str(n), b"pdf"
        def leer(nombre, contenido):
            eventos.append(f"lectura{nombre}")
            return {}, []
        with patch("asistencia_tecnica.analizar_solicitud_pdf", side_effect=leer):
            analizar_solicitudes(archivos())
        self.assertEqual(eventos, ["carga0", "lectura0", "carga1", "lectura1", "carga2", "lectura2"])


class NombresDescargaTest(unittest.TestCase):
    def test_zip_con_30_solicitudes_se_extrae_y_conserva_todos_los_sipse(self):
        numeros = (
            "165569 165422 165421 165420 165419 165418 165417 165416 165415 "
            "164857 164748 164493 164420 164354 164348 164290 164228 163607 "
            "163604 163601 163599 163594 163543 163536 163442 163440 163397 "
            "163054 163043 161356"
        ).split()
        filas = [solicitud(sipse=n, localidad="San Cristóbal", prorroga_solicitada="84") for n in numeros]
        contenido, _, resumen = generar_zip_asistencia(filas, PLANTILLA, FECHA_RESPUESTA, "Prueba")
        with TemporaryDirectory(prefix="plan-zip-") as carpeta:
            with ZipFile(BytesIO(contenido)) as z:
                self.assertEqual(len(z.namelist()), 3)
                for nombre in z.namelist():
                    self.assertLessEqual(len(nombre.encode("utf-8")), 180)
                z.extractall(carpeta)
            word = Path(carpeta) / resumen[0]["archivo"]
            self.assertEqual(word.name, "Alcaldía Local de San Cristóbal - SIPSE - 165569-161356 - 30 solicitudes.docx")
            doc = Document(word)
            solicitudes = [p.text for p in doc.paragraphs if p.text.startswith("SOLICITUD ")]
            self.assertEqual(solicitudes, [f"SOLICITUD {n}" for n in numeros])
            self.assertEqual(len(doc.tables), 30)
            self.assertEqual(doc.tables[-1].rows[9].cells[1].text, "84 días")

    def test_nombres_cortos_conservan_sipse_y_lotes_grandes_no_exceden_bytes(self):
        filas = [solicitud(sipse="160428"), solicitud(sipse="160427")]
        self.assertEqual(_nombre_word_localidad("Sumapaz", iter(filas)),
                         "Alcaldía Local de Sumapaz - SIPSE - 160428-160427 - 2 solicitudes.docx")
        self.assertEqual(_nombre_excel_asistencia(iter(filas)),
                         "Matriz asistencia técnica CPS - 160428-160427 - 2 solicitudes.xlsx")
        filas.insert(1, solicitud(sipse="160425"))
        self.assertEqual(_nombre_word_localidad("Sumapaz", filas),
                         "Alcaldía Local de Sumapaz - SIPSE - 160428-160427 - 3 solicitudes.docx")
        self.assertEqual(_nombre_excel_asistencia(filas),
                         "Matriz asistencia técnica CPS - 160428-160427 - 3 solicitudes.xlsx")
        self.assertEqual(_nombre_word_localidad("Sumapaz", filas[:1]),
                         "Alcaldía Local de Sumapaz - SIPSE - 160428 - 1 solicitud.docx")
        for filas in ([solicitud(sipse=str(164700+n)) for n in range(120)],
                      [solicitud(sipse="", archivo="", contrato="á" + "1"*300)]):
            for nombre in (_nombre_word_localidad("San Cristóbal", filas), _nombre_excel_asistencia(filas)):
                self.assertLessEqual(len(nombre.encode("utf-8")), 180)
                self.assertTrue(nombre.endswith((".docx", ".xlsx")))


if __name__ == "__main__":
    unittest.main()
