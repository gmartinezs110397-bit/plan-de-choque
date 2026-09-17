from datetime import datetime
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from cxp_cruce import (
    aplicar_desempate_en_contratos,
    procesar_localidad_cxp,
)
from hoja_estrategias import actualizar_hoja_estrategias
from hoja_suspendidos import (
    _aplicar_tope_mes_anterior,
    _fila_totales_seguimiento,
    _resolver_columnas_mes_seguimiento,
)
from hoja_liquidados_con_saldo import _fila_conteo_liquidados


FECHA = datetime(2026, 5, 29)
PROXIMOS = "Pr\u00f3ximos a perder"
TRAMITES = "Tr\u00e1mites sectores"
LOCALIDADES = ("San Crist\u00f3bal", "La Candelaria")
ENCABEZADOS = [
    "NOMBRE CONTRATISTA", "CLASIFICACI\u00d3N", "No. de Cto",
    "A\u00d1O SUSCRIPCI\u00d3N", "FECHA INICIAL", "FECHA FINALIZACION",
    "APROPIACION DISPONIBLE", "GIROS", "MONTO LIBERACIONES", "SALDO FINAL",
    "ESTADO ACTUAL", "RESPONSABLE",
]


def plantilla_localidad(localidad, reporte_anterior=0, fila_header=1):
    wb = Workbook()
    wb.remove(wb.active)
    matriz = []
    grupos = {
        "Suspendidos": [(100, "SUSPENDIDO"), (200, "SUSPENDIDO"), (0, "LIQUIDADO")],
        PROXIMOS: [(1000, "TERMINADO EN PROCESO DE LIQUIDACI\u00d3N")] * 29
        + [(0, "TERMINADO EN PROCESO DE LIQUIDACI\u00d3N"), (500, "LIQUIDADO")],
        TRAMITES: [(700, "EN EJECUCI\u00d3N"), (0, "EN EJECUCI\u00d3N")],
        "Cps por depurar": [(10, "TERMINADO"), (20, "TERMINADO"), (30, "TERMINADO"), (0, "LIQUIDADO")],
        "Liquidados con saldo": [(400, "LIQUIDADO"), (600, "LIQUIDADO"), (0, "LIQUIDADO")],
    }
    for nombre_hoja, contratos in grupos.items():
        ws = wb.create_sheet(nombre_hoja)
        es_cps = nombre_hoja == "Cps por depurar"
        es_par = nombre_hoja in ("Suspendidos", PROXIMOS, TRAMITES)
        header = 3 if es_cps else fila_header
        titulos = ENCABEZADOS + ["SALDO A 30 DE ABRIL"]
        if es_par:
            titulos.append("ESTADO A 30 DE ABRIL")
        for col, texto in enumerate(titulos, 1):
            ws.cell(header, col, texto)
            if header == 1:
                ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        inicio = 4 if header == 3 else 3
        for i, (saldo, estado) in enumerate(contratos):
            contratista = f"Contratista {nombre_hoja} {i}"
            valores = [contratista, "Obra", i + 1, 2025, None, None, 10000, 0, 0, saldo, estado, "Responsable", saldo]
            if es_par:
                valores.append(estado)
            if reporte_anterior is None:
                valores[12:] = [None] * (len(valores) - 12)
            for col, valor in enumerate(valores, 1):
                ws.cell(inicio + i, col, valor)
            matriz.append({
                "Localidad": f"Fondo de Desarrollo Local de {localidad}",
                "NOMBRE CONTRATISTA": contratista, "N\u00famero Contrato": i + 1,
                "A\u00f1o Suscripci\u00f3n": 2025, "Apropiaci\u00f3n": 10000,
                "Saldo Final": saldo, "Estado Actual": estado,
            })
        pie = inicio + len(contratos)
        if es_par:
            ws.cell(pie, 13, sum(s for s, _ in contratos))
            ws.cell(pie, 14, reporte_anterior)
        elif es_cps:
            ws.cell(1, 13, reporte_anterior)
            ws.cell(2, 13, sum(s for s, _ in contratos))
        else:
            ws.cell(pie, 13, sum(s for s, _ in contratos))
            ws.cell(pie + 1, 13, reporte_anterior)
    ws = wb.create_sheet("Estrategias")
    ws.merge_cells("B2:H2")
    ws["B2"] = "ESTRATEGIAS PLAN DE CHOQUE"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B2"].alignment = Alignment(horizontal="center")
    ws["B2"].fill = PatternFill("solid", fgColor="FF0000")
    for i, nombre in enumerate(grupos, 4):
        ws.cell(i, 2, nombre)
        ws.cell(i, 5, 0)
        ws.cell(i, 6, 0)
        ws.cell(i, 7, f"=C{i}-E{i}")
    ws["B9"] = "Total"
    ws["E9"] = "=SUM(E4:E8)"
    ws["F9"] = "=SUM(F4:F8)"
    out = BytesIO()
    wb.save(out)
    return out.getvalue(), pd.DataFrame(matriz)


class ConteosPlanChoqueTest(unittest.TestCase):
    def test_sin_conteo_anterior_no_se_anulan_los_contratos_actuales(self):
        for anterior in (None, 0, 0.0):
            with self.subTest(anterior=anterior):
                self.assertEqual(_aplicar_tope_mes_anterior(29, anterior), (29, 29))

    def test_conserva_tope_anterior_positivo_y_conteo_cero_real(self):
        self.assertEqual(_aplicar_tope_mes_anterior(29, 30), (29, 29))
        self.assertEqual(_aplicar_tope_mes_anterior(29, 28), (28, 29))
        self.assertEqual(_aplicar_tope_mes_anterior(0, 0), (0, 0))

    def test_cruce_individual_con_totales_anteriores_cero_o_vacios(self):
        for localidad in LOCALIDADES:
            for reporte in (0, None):
                for header in (1, 3):
                    with self.subTest(localidad=localidad, reporte=reporte, header=header):
                        datos, matriz = plantilla_localidad(localidad, reporte, header)
                        resultado = procesar_localidad_cxp(datos, matriz, localidad, FECHA)
                        self.assertTrue(all(
                            mensaje.startswith("Cps por depurar: no se pudieron aplicar los colores")
                            for mensaje in resultado["observaciones"]
                        ), resultado["observaciones"])
                        wb = load_workbook(BytesIO(resultado["bytes_contratos"]))
                        esperados = {"Suspendidos": 2, PROXIMOS: 29, TRAMITES: 1}
                        for nombre, cantidad in esperados.items():
                            ws = wb[nombre]
                            _, col_estado = _resolver_columnas_mes_seguimiento(ws, FECHA)
                            self.assertEqual(ws.cell(_fila_totales_seguimiento(ws), col_estado).value, cantidad)
                        ws = wb["Liquidados con saldo"]
                        self.assertEqual(ws.cell(_fila_conteo_liquidados(ws), 14).value, 2)
                        self.assertEqual(wb["Cps por depurar"].cell(1, 14).value, 3)
                        ws = wb["Estrategias"]
                        self.assertEqual([ws.cell(f, 5).value for f in range(4, 9)], [2, 29, 1, 3, 2])
                        self.assertEqual([ws.cell(f, 6).value for f in range(4, 9)], [300, 29500, 700, 60, 1000])
                        self.assertEqual(ws["B2"].value, f"ESTRATEGIAS PLAN DE CHOQUE {localidad.upper()} 2026")
                        self.assertEqual(str(ws.merged_cells), "B2:H2")
                        self.assertEqual(ws["B2"].font.size, 14)
                        self.assertEqual(ws["B2"].fill.fgColor.rgb, "00FF0000")
                        self.assertEqual(ws["E9"].value, "=SUM(E4:E8)")
                        self.assertEqual(ws["G5"].value, "=C5-E5")
                        self.assertEqual(wb[PROXIMOS].cell(header, 13).value, "Saldo a 30 de abril")
                        self.assertEqual(wb[PROXIMOS].cell(_fila_totales_seguimiento(wb[PROXIMOS]), 14).value, reporte)

    def test_titulo_cambia_localidad_y_ano_conservando_formato(self):
        datos, _ = plantilla_localidad("Kennedy")
        wb = load_workbook(BytesIO(datos))
        ws = wb["Estrategias"]
        ws["B2"] = "ESTRATEGIAS PLAN DE CHOQUE KENNEDY 2026"
        actualizar_hoja_estrategias(ws, datetime(2027, 1, 15), {}, "La Candelaria")
        self.assertEqual(ws["B2"].value, "ESTRATEGIAS PLAN DE CHOQUE LA CANDELARIA 2027")
        self.assertEqual(ws["B2"].font.size, 14)

    def test_desempate_tambien_actualiza_titulo(self):
        datos, _ = plantilla_localidad("San Crist\u00f3bal")
        nuevos, _ = aplicar_desempate_en_contratos(datos, FECHA, {}, [], "San Crist\u00f3bal")
        wb = load_workbook(BytesIO(nuevos))
        self.assertEqual(wb["Estrategias"]["B2"].value, "ESTRATEGIAS PLAN DE CHOQUE SAN CRIST\u00d3BAL 2026")

    def test_llamadas_sin_localidad_conservan_el_titulo(self):
        datos, _ = plantilla_localidad("Kennedy")
        ws = load_workbook(BytesIO(datos))["Estrategias"]
        titulo = ws["B2"].value
        actualizar_hoja_estrategias(ws, FECHA, {})
        self.assertEqual(ws["B2"].value, titulo)


if __name__ == "__main__":
    unittest.main()
