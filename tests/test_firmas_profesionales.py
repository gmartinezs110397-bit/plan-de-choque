from datetime import date
from io import BytesIO
from pathlib import Path
import unittest
from docx import Document
from docx.oxml.ns import qn
from asistencia_tecnica import generar_documento_localidad
from test_asistencia_tecnica_word import PLANTILLA, solicitud


class FirmasProfesionalesTest(unittest.TestCase):
    def documento(self, profesional):
        return Document(BytesIO(generar_documento_localidad([solicitud()], PLANTILLA, date(2026,9,28), profesional)))

    def test_firma_correcta_junto_a_cada_nombre(self):
        doc = self.documento('Ingrith Khaterine Martínez Sánchez')
        for etiqueta, nombre, archivo in (
            ('Proyectó:', 'Ingrith Khaterine Martínez Sánchez', 'ingrith_martinez.png'),
            ('Revisó:', 'Andres González', 'andres_gonzalez.jpeg'),
        ):
            p = next(p for p in doc.paragraphs if p.text.startswith(etiqueta))
            self.assertEqual(p.text.strip(), f'{etiqueta} {nombre} - Profesional DGDL')
            dibujos = p._p.xpath('.//w:drawing')
            self.assertEqual(len(dibujos), 1)
            rid = dibujos[0].xpath('.//a:blip')[0].get(qn('r:embed'))
            self.assertEqual(doc.part.rels[rid].target_part.blob, (PLANTILLA.parent / 'firmas' / archivo).read_bytes())

    def test_otro_profesional_no_recibe_firma_de_ingrith(self):
        doc = self.documento('OTRO PROFESIONAL')
        p = next(p for p in doc.paragraphs if p.text.startswith('Proyectó:'))
        self.assertEqual(len(p._p.xpath('.//w:drawing')), 0)

    def test_conserva_recorte_original_del_revisor(self):
        doc = self.documento('Ingrith Khaterine Martínez Sánchez')
        p = next(p for p in doc.paragraphs if p.text.startswith('Revisó:'))
        rect = p._p.xpath('.//a:srcRect')[0]
        self.assertEqual(dict(rect.attrib), {'l':'5185','t':'24011','r':'5556','b':'29379'})
