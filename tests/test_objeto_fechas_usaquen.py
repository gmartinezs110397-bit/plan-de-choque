from datetime import date
from io import BytesIO
import unittest
from docx import Document
from asistencia_tecnica import _limpiar_objeto_solicitud, parsear_fecha, _calcular_validacion_solicitud
from test_asistencia_tecnica_word import generar, solicitud


class ObjetoFechasUsaquenTest(unittest.TestCase):
    def test_pagina_web_forma_parte_del_objeto(self):
        objeto='GENERAR CONTENIDOS PARA LAS REDES SOCIALES Y LA PÁGINA WEB DE LA ALCALDÍA LOCAL.'
        self.assertEqual(_limpiar_objeto_solicitud(objeto), objeto)
        self.assertEqual(_limpiar_objeto_solicitud(objeto+' Página 1 de 3'), objeto)

    def test_fecha_con_espacios_entre_componentes(self):
        for texto in ('02/ 02/ 2026', '02 /02 / 2026', '02 / 02 / 2026', '02- 02 -2026', '02/\n02/2026'):
            with self.subTest(texto=texto):
                self.assertEqual(parsear_fecha(texto),date(2026,2,2))
        self.assertIsNone(parsear_fecha('31/ 02/ 2026'))
        self.assertIsNone(parsear_fecha('fecha desconocida'))

    def test_validacion_con_fechas_espaciadas(self):
        fila=self.fila()
        validacion=_calcular_validacion_solicitud(fila)
        for campo in ('se_ajusta_fecha_fin','se_ajusta_prorroga','se_ajusta_fecha_final','se_ajusta_adicion'):
            self.assertTrue(validacion[campo],campo)

    def fila(self):
        return solicitud(
            fecha_inicio='02/ 02/ 2026',fecha_terminacion_inicial='01/ 08/ 2026',
            fecha_terminacion_final='01/11/2026',plazo_inicial='6 meses',
            prorroga_solicitada='3 meses',valor_inicial=39000000,valor_adicion=19500000,
            objeto='CONTENIDOS PARA LA PÁGINA WEB DE LA ALCALDÍA LOCAL.',
        )

    def test_word_conserva_objeto_y_sin_observaciones(self):
        doc=Document(BytesIO(generar([self.fila()])))
        self.assertIn('PÁGINA WEB DE LA ALCALDÍA LOCAL.',doc.tables[0].rows[4].cells[1].text)
        textos=[p.text for p in doc.paragraphs]
        self.assertIn('• Sin observaciones',textos)
        self.assertFalse(any('faltan datos' in t for t in textos))
