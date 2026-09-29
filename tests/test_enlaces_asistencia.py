from io import BytesIO
from zipfile import ZipFile
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook

from enlaces_asistencia import leer_base_enlaces, cruzar_enlaces, generar_excel_enlaces, NOMBRE_EXCEL_ENLACES
from asistencia_tecnica import generar_zip_asistencia
from test_asistencia_tecnica_word import solicitud, PLANTILLA, FECHA_RESPUESTA


URL = "https://community.secop.gov.co/Public/Tendering/OpportunityDetail/Index?noticeUID=CO1.NTC.9775874&isModal=true"


def base(registros):
    wb = Workbook()
    ws = wb.active
    encabezados = [None] * 49
    for pos, texto in [(0, "Nombre Entidad"), (10, "ID Contrato"), (11, "Número del Contrato"), (48, "URLProceso")]:
        encabezados[pos] = texto
    ws.append(encabezados)
    for entidad, numero, identificador, url in registros:
        fila = [None] * 49
        for pos, valor in [(0, entidad), (10, identificador), (11, numero), (48, url)]:
            fila[pos] = valor
        ws.append(fila)
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class EnlacesAsistenciaTest(unittest.TestCase):
    def setUp(self):
        self.fila = solicitud(localidad="Usaquén", contrato="290-2026CPS-P-(152221)", sipse="156144")
        self.registro = ("ALCALDIA LOCAL DE USAQUEN", "290-2026 CPS-P-(152221)", "CO1.PCCNTR.1", URL)

    def test_columnas_a_l_aw_y_tolerancia_espacios_sin_cambiar_datos(self):
        indice = leer_base_enlaces(base([self.registro]))
        fila = cruzar_enlaces([self.fila], indice)[0]
        self.assertEqual(fila["link_secop"], URL)
        self.assertEqual(fila["entidad_enlace"], self.registro[0])
        for campo in self.fila:
            self.assertEqual(fila[campo], self.fila[campo])

    def test_no_cruza_otro_numero_alcaldia_tipo_o_identificador(self):
        indice = leer_base_enlaces(base([self.registro]))
        for cambios in [dict(localidad="Suba"), dict(contrato="291-2026 CPS-P-(152221)"),
                        dict(contrato="290-2026 CPS-AG-(152221)"), dict(contrato="290-2026 CPS-P-(152222)")]:
            fila = cruzar_enlaces([{**self.fila, **cambios}], indice)[0]
            self.assertEqual(fila["link_secop"], "")
            self.assertEqual(fila["estado_enlace"], "No encontrado en la base")

    def test_duplicado_identico_no_ambiguo_y_contratos_distintos_si(self):
        indice = leer_base_enlaces(base([self.registro, self.registro]))
        self.assertEqual(cruzar_enlaces([self.fila], indice)[0]["link_secop"], URL)
        distinto = (*self.registro[:2], "CO1.PCCNTR.2", URL)
        indice = leer_base_enlaces(base([self.registro, distinto]))
        fila = cruzar_enlaces([self.fila], indice)[0]
        self.assertEqual(fila["link_secop"], "")
        self.assertIn("varias coincidencias", fila["estado_enlace"])

    def test_no_base_vacia_y_url_no_valida_no_inventan_enlace(self):
        for indice in [None, {}, leer_base_enlaces(base([(*self.registro[:3], "javascript:alert(1)")]))]:
            fila = cruzar_enlaces([{**self.fila, "link_secop": URL}], indice)[0]
            self.assertEqual(fila["link_secop"], "")
            self.assertNotEqual(fila["estado_enlace"], "Encontrado")

    def test_encabezados_invalidos_explican_error(self):
        wb = Workbook()
        salida = BytesIO()
        wb.save(salida)
        with self.assertRaisesRegex(ValueError, "URLProceso"):
            leer_base_enlaces(salida.getvalue())

    def test_excel_y_word_mismo_orden_reinician_numeracion_por_localidad(self):
        filas = [solicitud(localidad="Suba", sipse="164702", contrato="B"),
                 self.fila, solicitud(localidad="Suba", sipse="164719", contrato="A")]
        filas = cruzar_enlaces(filas, leer_base_enlaces(base([self.registro])))
        contenido, _, resumen = generar_zip_asistencia(filas, PLANTILLA, FECHA_RESPUESTA, "Prueba")
        with ZipFile(BytesIO(contenido)) as z:
            ws = load_workbook(BytesIO(z.read(NOMBRE_EXCEL_ENLACES))).active
            datos = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual([(r[0], r[1], r[2]) for r in datos],
                             [("Usaquén", 1, "156144"), ("Suba", 1, "164719"), ("Suba", 2, "164702")])
            self.assertEqual(ws["F2"].hyperlink.target, URL)
            self.assertIsNone(ws["F3"].hyperlink)
            self.assertEqual(ws["E2"].value, self.registro[0])
            for item in resumen:
                doc = Document(BytesIO(z.read(item["archivo"])))
                sipses = [p.text.split()[-1] for p in doc.paragraphs if p.text.strip().startswith("SOLICITUD ")]
                self.assertEqual(sipses, [r[2] for r in datos if r[0] == item["localidad"]])

    def test_contenido_es_texto_no_formula_y_preserva_duplicadas(self):
        filas = [solicitud(contrato="=1+1", entidad_enlace="=1+1")] * 2
        ws = load_workbook(BytesIO(generar_excel_enlaces(filas))).active
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws["D2"].data_type, "s")
        self.assertEqual(ws["E2"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
