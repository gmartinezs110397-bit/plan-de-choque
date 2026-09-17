import ast
from dataclasses import asdict, replace
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import pickle
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

import pandas as pd
from openpyxl import load_workbook

from tabla_resumen_proyecto import (
    FilaDepuraTipoCto,
    FilaDismAntiguedadCto,
    crear_excel_tabla_resumen_proyecto,
    fila_depurados_desde_matriz,
    filas_depura_tipos_ctos_desde_matriz,
    filas_dism_antiguedad_ctos_desde_matriz,
    fila_por_depurar_vigencia_desde_matriz,
    fila_cps_pn_desde_matriz,
    filas_depurar_x_vigencia_desde_matriz,
)


class DepuraTiposContratosTest(unittest.TestCase):
    def matriz(self):
        return pd.DataFrame({
            "Clasificación": [
                " Obra Pública ", "OBRA PUBLICA", "Acta", "Acta", "Acta",
                None, "SUMINISTRO", None,
            ],
            "Año Suscripción": [2024, 2024, 2025, 2025, 2025, 2023, None, None],
            "Saldo Final": [100, 0, 0, "0", -10, 15, 50, 180],
            "Estado Actual": ["En Ejecución", "Liquidado", "Liquidado", "Liquidado",
                              "Suspendido", "Terminado En Proceso De Liquidación", "Liquidado", None],
        })

    def test_agrupa_clasificaciones_y_excluye_total_sin_vigencia(self):
        filas = filas_depura_tipos_ctos_desde_matriz(self.matriz(), "Usaquén")
        conteos = {f.tipo_contrato: (f.cantidad_inicial, f.depurados) for f in filas}
        self.assertEqual(conteos, {
            "Acta": (3, 2), "Obra Pública": (2, 1), "Sin clasificación": (1, 0),
        })
        total = fila_depurados_desde_matriz(self.matriz(), "Usaquén")
        self.assertEqual(sum(f.cantidad_inicial for f in filas), total.cantidad_inicial)
        self.assertEqual(sum(f.depurados for f in filas), total.depurados)

    def test_tipos_cps_persona_natural_y_juridica_se_mantienen_separados(self):
        df = pd.DataFrame({
            "Clasificación": [
                "Prestación de Servicios (Persona Natural)",
                "Prestación de Servicios (Persona Jurídica)",
            ],
            "Año Suscripción": [2025, 2025], "Saldo Final": [0, 50],
        })
        filas = filas_depura_tipos_ctos_desde_matriz(df, "Kennedy")
        self.assertEqual(len(filas), 2)
        self.assertEqual({f.depurados for f in filas}, {0, 1})

    def test_columnas_faltantes_se_informan(self):
        with self.assertRaisesRegex(ValueError, "Clasificación"):
            filas_depura_tipos_ctos_desde_matriz(
                self.matriz().drop(columns="Clasificación"), "Kennedy"
            )

    def test_exporta_formulas_fecha_y_tablas_por_localidad(self):
        filas = filas_depura_tipos_ctos_desde_matriz(self.matriz(), "Usaquén")
        filas.append(FilaDepuraTipoCto(8, "KENNEDY", "Acta", 10, 4))
        datos = json.loads(json.dumps([asdict(f) for f in reversed(filas)]))
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), filas_depura_tipos_ctos=datos,
        )))
        self.assertEqual(wb.sheetnames, [
            "LIB Y FEN", "CON PÉRDIDA", "PRÓXIMOS A PERDER", "BOGDATA VS MATRIZ",
            "DEPURADOS", "POR DEPURAR x VIGENCIA", "CPS PN", "DEPURA. TIPOS CTOS",
            "DISM. ANTIGÜEDAD CTOS",
            "DEPURAR X VIGENCIA",
        ])
        ws = wb["DEPURA. TIPOS CTOS"]
        self.assertEqual(ws["A1"].value, "USAQUÉN")
        self.assertEqual(ws["C2"].value, "CTOS DEPURADOS A 30 DE JUNIO DE 2026")
        self.assertEqual(ws["D3"].value, "=B3-C3")
        self.assertEqual(ws["E3"].value, "=IF(B3=0,0,C3/B3)")
        self.assertEqual(ws["B6"].value, "=SUM(B3:B5)")
        self.assertEqual(ws["C6"].value, "=SUM(C3:C5)")
        self.assertEqual(ws["D6"].value, "=B6-C6")
        self.assertEqual(ws["E6"].value, "=IF(B6=0,0,C6/B6)")
        self.assertEqual(ws["E3"].number_format, "0.00%")
        self.assertEqual(len(ws._charts), 2)
        kennedy_row = next(c.row for c in ws["A"] if c.value == "KENNEDY")
        self.assertGreater(kennedy_row, ws._charts[0].anchor._from.row)
        self.assertEqual(ws.cell(kennedy_row + 2, 2).value, 10)
        self.assertEqual(ws.cell(kennedy_row + 2, 3).value, 4)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "00C00000")
        self.assertEqual(ws["A2"].fill.fgColor.rgb, "00FFC000")
        wb.close()

    def test_corte_cambia_con_el_mes_y_el_anio(self):
        filas = [FilaDepuraTipoCto(1, "USAQUÉN", "Acta", 1, 1)]
        for corte, texto in [
            (date(2028, 2, 29), "29 DE FEBRERO DE 2028"),
            (date(2026, 12, 31), "31 DE DICIEMBRE DE 2026"),
        ]:
            with self.subTest(corte=corte):
                wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
                    [], [], corte, filas_depura_tipos_ctos=filas,
                )))
                self.assertIn(texto, wb["DEPURA. TIPOS CTOS"]["C2"].value)
                wb.close()

    def test_sin_datos_no_inventa_contratos(self):
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 5, 31),
        )))
        ws = wb["DEPURA. TIPOS CTOS"]
        self.assertIn("Sin datos", ws["A1"].value)
        self.assertEqual(len(ws._charts), 0)
        wb.close()

    def test_clasificacion_que_empieza_con_igual_es_texto(self):
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 5, 31),
            filas_depura_tipos_ctos=[FilaDepuraTipoCto(1, "USAQUÉN", "=1+1", 1, 0)],
        )))
        self.assertEqual(wb["DEPURA. TIPOS CTOS"]["A3"].data_type, "s")
        wb.close()

    def test_conteos_se_conservan_al_externalizar_y_recargar_sesion(self):
        class Sesion(dict):
            def __setattr__(self, nombre, valor):
                self[nombre] = valor

        filas = [asdict(f) for f in filas_depura_tipos_ctos_desde_matriz(self.matriz(), "Usaquén")]
        antiguedad = [asdict(f) for f in filas_dism_antiguedad_ctos_desde_matriz(self.matriz(), "Usaquén")]
        pendientes = [asdict(f) for f in filas_depurar_x_vigencia_desde_matriz(self.matriz(), "Usaquén")]
        sesion = Sesion(
            cruce_informe=[{"localidad": "Usaquén"}], tabla_resumen_depura_tipos_ctos=filas,
            tabla_resumen_dism_antiguedad_ctos=antiguedad,
            tabla_resumen_depurar_x_vigencia=pendientes,
        )
        nombres = {
            "_persistir_snapshot_consolidacion", "_cargar_snapshot_consolidacion_cache",
            "_cargar_snapshot_consolidacion", "_tabla_resumen_depura_tipos_ctos_para_ui",
            "_tabla_resumen_dism_antiguedad_ctos_para_ui",
            "_tabla_resumen_depurar_x_vigencia_para_ui",
        }
        arbol = ast.parse(Path("app.py").read_text(encoding="utf-8"))
        funciones = [n for n in arbol.body if isinstance(n, ast.FunctionDef) and n.name in nombres]
        for funcion in funciones:
            funcion.decorator_list = []
        with TemporaryDirectory() as carpeta:
            contexto = {
                "st": SimpleNamespace(session_state=sesion), "Path": Path, "pickle": pickle,
                "APP_SESSION_VERSION": "build-prueba",
                "_CLAVE_SNAPSHOT": "snapshot", "_CLAVE_RESUMEN_LIGERO": "resumen",
                "_directorio_archivos_sesion": lambda: Path(carpeta),
                "_resumen_ligero_desde_informe": lambda informe: {},
            }
            exec(compile(ast.Module(body=funciones, type_ignores=[]), "app.py", "exec"), contexto)
            contexto["_persistir_snapshot_consolidacion"]()
            self.assertEqual(sesion["tabla_resumen_depura_tipos_ctos"], [])
            self.assertEqual(contexto["_tabla_resumen_depura_tipos_ctos_para_ui"](), filas)
            self.assertEqual(sesion["tabla_resumen_dism_antiguedad_ctos"], [])
            self.assertEqual(contexto["_tabla_resumen_dism_antiguedad_ctos_para_ui"](), antiguedad)
            self.assertEqual(sesion["tabla_resumen_depurar_x_vigencia"], [])
            self.assertEqual(contexto["_tabla_resumen_depurar_x_vigencia_para_ui"](), pendientes)

    def test_descarga_global_incluye_la_nueva_pestana(self):
        filas = [asdict(f) for f in filas_depura_tipos_ctos_desde_matriz(self.matriz(), "Usaquén")]
        arbol = ast.parse(Path("app.py").read_text(encoding="utf-8"))
        funcion = next(n for n in arbol.body if isinstance(n, ast.FunctionDef)
                       and n.name == "construir_archivos_salida_global")
        contexto = {
            "pd": pd, "st": SimpleNamespace(session_state={}),
            "fecha_referencia_analisis": lambda: date(2026, 6, 30),
            "localidades_analizadas": lambda stats: ["Usaquén"],
            "construir_avance_plan_de_choque": lambda *args: pd.DataFrame(),
            "generate_excel_bytes": lambda *args, **kwargs: BytesIO(b"avance"),
            "nombre_archivo_salida": lambda base, *args: base + ".xlsx",
            "ARCHIVO_AVANCE_BASE": "Avance plan de choque",
            "ARCHIVO_RESUMEN_BASE": "Tabla Resumen Proyecto",
            "_tabla_resumen_depura_tipos_ctos_para_ui": lambda: filas,
            "_tabla_resumen_dism_antiguedad_ctos_para_ui": lambda: [
                asdict(f) for f in filas_dism_antiguedad_ctos_desde_matriz(self.matriz(), "Usaquén")
            ],
            "_tabla_resumen_depurar_x_vigencia_para_ui": lambda: [
                asdict(f) for f in filas_depurar_x_vigencia_desde_matriz(self.matriz(), "Usaquén")
            ],
        }
        for sufijo in ("lib_y_fen", "con_perdida", "proximos_a_perder", "bogdata_matriz",
                       "depurados", "por_depurar_vigencia", "cps_pn"):
            contexto[f"_tabla_resumen_{sufijo}_para_ui"] = lambda: []
        exec(compile(ast.Module(body=[funcion], type_ignores=[]), "app.py", "exec"), contexto)
        archivos = dict(contexto["construir_archivos_salida_global"](pd.DataFrame(), [{}]))
        wb = load_workbook(BytesIO(archivos["Tabla Resumen Proyecto.xlsx"]))
        self.assertEqual(wb["DEPURA. TIPOS CTOS"]["B3"].value, 3)
        self.assertEqual(wb["DEPURA. TIPOS CTOS"]["C3"].value, 2)
        self.assertEqual(wb["DISM. ANTIGÜEDAD CTOS"]["A3"].value, 2023)
        self.assertEqual(wb["DISM. ANTIGÜEDAD CTOS"]["B3"].value, 1)
        self.assertEqual(wb["DISM. ANTIGÜEDAD CTOS"]["C3"].value, 0)
        self.assertEqual(wb["DEPURAR X VIGENCIA"]["A3"].value, "CONTRATOS POR DEPURAR POR VIGENCIAS")
        wb.close()


class DismAntiguedadContratosTest(unittest.TestCase):
    def matriz(self):
        return pd.DataFrame({
            "Año de Suscripción": [2025, "2022", 2023.0, 2024, 2025, 2023, 2024, 2024, 2026, 2010, None],
            "Saldo Final": [0, "0", 100, 0, 50, -5, 2, 3, 1, 0, 9999],
        })

    def test_agrupa_todos_los_anos_y_solo_depura_saldos_cero(self):
        filas = filas_dism_antiguedad_ctos_desde_matriz(self.matriz(), "Usaquén")
        self.assertEqual([(f.vigencia, f.cantidad_inicial, f.depurados) for f in filas], [
            (2010, 1, 1), (2022, 1, 1), (2023, 2, 0), (2024, 3, 1), (2025, 2, 1), (2026, 1, 0),
        ])
        self.assertEqual(sum(f.cantidad_inicial for f in filas), 10)
        self.assertEqual(sum(f.depurados for f in filas), 4)
        self.assertTrue(all(f.numero == 1 and f.localidad == "USAQUÉN" for f in filas))

    def test_ano_no_numerico_no_se_omite_ni_se_trunca(self):
        df = pd.DataFrame({
            "Año Suscripción": ["Sin info", 2024.5, 2024], "Saldo Final": [0, 10, 0],
        })
        filas = filas_dism_antiguedad_ctos_desde_matriz(df, "Kennedy")
        self.assertEqual([(f.vigencia, f.cantidad_inicial, f.depurados) for f in filas], [
            (2024, 1, 1), (None, 2, 1),
        ])

    def test_exporta_anos_totales_y_formulas_por_localidad(self):
        filas = filas_dism_antiguedad_ctos_desde_matriz(self.matriz(), "Usaquén")
        filas.append(FilaDismAntiguedadCto(8, "KENNEDY", 2021, 9, 5))
        datos = json.loads(json.dumps([asdict(f) for f in reversed(filas)]))
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), filas_dism_antiguedad_ctos=datos,
        )))
        ws = wb["DISM. ANTIGÜEDAD CTOS"]
        self.assertEqual(ws["A1"].value, "USAQUÉN")
        self.assertEqual(ws["C2"].value, "CTOS LIQUIDADOS/\nLIBERADOS/FENECIDOS\nA 30 DE JUNIO DE 2026")
        self.assertEqual([ws.cell(r, 1).value for r in range(3, 9)], [2010, 2022, 2023, 2024, 2025, 2026])
        self.assertEqual(ws["A3"].number_format, "0")
        self.assertEqual(ws["D3"].value, "=B3-C3")
        self.assertEqual(ws["A9"].value, "TOTAL")
        self.assertEqual(ws["B9"].value, "=SUM(B3:B8)")
        self.assertEqual(ws["C9"].value, "=SUM(C3:C8)")
        self.assertEqual(ws["D9"].value, "=B9-C9")
        self.assertEqual(ws["A12"].value, "KENNEDY")
        self.assertEqual(ws["A14"].value, 2021)
        self.assertEqual(ws["B14"].value, 9)
        self.assertEqual(ws["C14"].value, 5)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "00C00000")
        self.assertEqual(ws["A2"].fill.fgColor.rgb, "00FFC000")
        self.assertEqual(ws.max_column, 4)
        wb.close()

    def test_sin_datos_muestra_ausencia_sin_inventar_contratos(self):
        filas = filas_dism_antiguedad_ctos_desde_matriz(self.matriz().iloc[:0], "Kennedy")
        self.assertEqual(filas, [])
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), filas_dism_antiguedad_ctos=filas,
        )))
        self.assertIn("Sin datos", wb["DISM. ANTIGÜEDAD CTOS"]["A1"].value)
        wb.close()

    def test_columnas_faltantes_se_informan(self):
        with self.assertRaisesRegex(ValueError, "Saldo Final"):
            filas_dism_antiguedad_ctos_desde_matriz(self.matriz().drop(columns="Saldo Final"), "Kennedy")


class DepurarXVigenciaTest(unittest.TestCase):
    def matriz(self):
        return pd.DataFrame({
            "NOMBRE CONTRATISTA": ["Uno", "Dos", "Tres", "Cuatro", "Cinco", "Seis", "Siete", "Ocho", "Nueve", None],
            "Año Suscripción": [2023, 2023, 2024, 2024, 2024, 2025, None, 2010, 2026, None],
            "Clasificación": ["Acta"] * 6 + [None] + ["Acta"] * 2 + [None],
            "Saldo Final": [10, 20, 100, 0, 30, -5, 15, 7, 8, 185],
            "Estado Actual": ["en ejecución", " EN EJECUCIÓN ", "Suspendido", "Liquidado", "Liquidado",
                              "TERMINADO (no se liquida)", None, "Terminado En Proceso De Liquidación",
                              "En Revisión Por Entes De Control", None],
        })

    def parametros(self, df=None, localidad="Usaquén"):
        if df is None:
            df = self.matriz()
        return {
            "filas_depurados": [fila_depurados_desde_matriz(df, localidad)],
            "filas_depura_tipos_ctos": filas_depura_tipos_ctos_desde_matriz(df, localidad),
            "filas_dism_antiguedad_ctos": filas_dism_antiguedad_ctos_desde_matriz(df, localidad),
            "filas_por_depurar_vigencia": [fila_por_depurar_vigencia_desde_matriz(df, localidad)],
            "filas_depurar_x_vigencia": filas_depurar_x_vigencia_desde_matriz(df, localidad),
        }

    def test_cuenta_y_suma_por_estado_y_ano_solo_saldos_no_cero(self):
        filas = filas_depurar_x_vigencia_desde_matriz(self.matriz(), "Usaquén")
        datos = {(f.vigencia, f.estado): (f.cantidad, f.saldo_final) for f in filas}
        self.assertEqual(datos[(2023, "En Ejecución")], (2, 30))
        self.assertEqual(datos[(2024, "Liquidado")], (1, 30))
        self.assertEqual(datos[(2024, "Suspendido")], (1, 100))
        self.assertEqual(datos[(None, "Sin información")], (1, 15))
        self.assertEqual(datos[(2025, "TERMINADO (no se liquida)")], (1, -5))
        self.assertEqual(sum(f.cantidad for f in filas), 8)
        self.assertEqual(sum(f.saldo_final for f in filas), 185)

    def test_pendientes_coinciden_en_todos_los_resumenes_generales(self):
        parametros = self.parametros()
        total = parametros["filas_depurados"][0]
        self.assertEqual(total.cantidad_inicial, 9)
        self.assertEqual(total.cantidad_inicial - total.depurados, 8)
        for nombre in ("filas_depura_tipos_ctos", "filas_dism_antiguedad_ctos"):
            self.assertEqual(sum(f.cantidad_inicial - f.depurados for f in parametros[nombre]), 8)
        conteos = parametros["filas_por_depurar_vigencia"][0].conteos_por_vigencia
        self.assertEqual(sum(c[1] for c in conteos.values()), 8)
        self.assertEqual(conteos[None], (1, 1))
        self.assertEqual(conteos[2010], (1, 1))
        self.assertEqual(conteos[2026], (1, 1))

    def test_exporta_dos_tablas_por_localidad_y_estados_adicionales(self):
        parametros = self.parametros()
        parametros = {k: json.loads(json.dumps([asdict(f) for f in filas])) for k, filas in parametros.items()}
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), **parametros,
        )))
        ws = wb["DEPURAR X VIGENCIA"]
        self.assertEqual(ws["A1"].value, "USAQUÉN")
        self.assertEqual(ws["A3"].value, "CONTRATOS POR DEPURAR POR VIGENCIAS")
        self.assertEqual(ws["A14"].value, "SALDOS POR DEPURAR POR VIGENCIAS")
        self.assertEqual(ws["A10"].value, "Sin información")
        self.assertEqual(ws["B6"].value, 2)
        self.assertEqual(ws["D7"].value, 0)
        self.assertEqual(ws["E7"].value, 1)
        self.assertEqual(ws["I11"].value, "=SUM(I5:I10)")
        self.assertEqual(ws["B12"].value, "=IF($I$11=0,0,B11/$I$11)")
        self.assertEqual(ws["B12"].number_format, "0.00%")
        self.assertEqual(ws["C18"].value, 100)
        self.assertEqual(ws["I22"].value, "=SUM(I16:I21)")
        self.assertIn("$", ws["C18"].number_format)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "00C00000")
        self.assertEqual(ws["A4"].fill.fgColor.rgb, "00FFC000")
        antigua = wb["POR DEPURAR x VIGENCIA"]
        encabezados = [c.value for c in antigua[1]]
        self.assertIn(2010, encabezados)
        self.assertIn(2026, encabezados)
        col_sin_info = encabezados.index("Sin info") + 1
        self.assertEqual(antigua.cell(3, col_sin_info + 1).value, 1)
        wb.close()

    def test_impide_exportar_totales_distintos(self):
        parametros = self.parametros()
        parametros["filas_depurar_x_vigencia"] = parametros["filas_depurar_x_vigencia"][:-1]
        with self.assertRaisesRegex(ValueError, "No coinciden.*USAQUÉN"):
            crear_excel_tabla_resumen_proyecto([], [], date(2026, 6, 30), **parametros)

    def test_valida_por_localidad_aunque_el_total_global_coincida(self):
        parametros = self.parametros()
        kennedy = self.parametros(localidad="Kennedy")
        for nombre in parametros:
            parametros[nombre].extend(kennedy[nombre])
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), **parametros,
        )))
        valores = [c.value for c in wb["DEPURAR X VIGENCIA"]["A"]]
        self.assertIn("KENNEDY", valores)
        self.assertEqual(valores.count("CONTRATOS POR DEPURAR POR VIGENCIAS"), 2)
        self.assertEqual(valores.count("SALDOS POR DEPURAR POR VIGENCIAS"), 2)
        wb.close()
        parametros["filas_depurar_x_vigencia"] = [
            replace(f, numero=8, localidad="KENNEDY") for f in parametros["filas_depurar_x_vigencia"]
        ]
        self.assertEqual(sum(f.cantidad for f in parametros["filas_depurar_x_vigencia"]), 16)
        with self.assertRaisesRegex(ValueError, "No coinciden.*USAQUÉN"):
            crear_excel_tabla_resumen_proyecto([], [], date(2026, 6, 30), **parametros)

    def test_contratos_sin_pendientes_muestran_dos_tablas_con_totales_cero(self):
        df = self.matriz().copy()
        df["Saldo Final"] = 0
        wb = load_workbook(BytesIO(crear_excel_tabla_resumen_proyecto(
            [], [], date(2026, 6, 30), **self.parametros(df),
        )))
        ws = wb["DEPURAR X VIGENCIA"]
        self.assertEqual(ws["A3"].value, "CONTRATOS POR DEPURAR POR VIGENCIAS")
        self.assertEqual(ws["F5"].value, 0)
        self.assertEqual(ws["A8"].value, "SALDOS POR DEPURAR POR VIGENCIAS")
        self.assertEqual(ws["F10"].value, 0)
        wb.close()

    def test_saldos_faltantes_no_se_confunden_con_cero_en_ninguna_pestana(self):
        for valor in (None, "", "No informado", float("inf")):
            df = self.matriz().copy()
            df["Saldo Final"] = df["Saldo Final"].astype(object)
            df.loc[0, "Saldo Final"] = valor
            for funcion in (fila_depurados_desde_matriz, filas_depura_tipos_ctos_desde_matriz,
                            filas_dism_antiguedad_ctos_desde_matriz, fila_por_depurar_vigencia_desde_matriz,
                            filas_depurar_x_vigencia_desde_matriz):
                with self.subTest(valor=valor, funcion=funcion.__name__):
                    with self.assertRaisesRegex(ValueError, "Saldo Final vacío o no numérico"):
                        funcion(df, "Usaquén")

    def test_columnas_faltantes_se_informan(self):
        with self.assertRaisesRegex(ValueError, "Estado Actual"):
            filas_depurar_x_vigencia_desde_matriz(self.matriz().drop(columns="Estado Actual"), "Usaquén")


if __name__ == "__main__":
    unittest.main()
