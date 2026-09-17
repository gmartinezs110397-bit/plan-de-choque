from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

from asistencia_tecnica import generar_documento_localidad, generar_zip_asistencia


ROOT = Path(__file__).resolve().parents[1]
PLANTILLA = ROOT / "templates/asistencia_tecnica/MODELO_DE_RESPUESTA_CPS.docx"
FECHA_RESPUESTA = date(2026, 9, 17)
INTRODUCCION = (
    "La Direcci\u00f3n para la Gesti\u00f3n del Desarrollo Local (DGDL) en cumplimiento con lo "
    "dispuesto en la Constituci\u00f3n Pol\u00edtica, la Ley 80 de 1993, la Ley 1150 de 2007 en "
    "sus art\u00edculos 2, 3, y 4 y sus decretos reglamentarios, el art\u00edculo 15 del Decreto "
    "Distrital 642 de 2025, as\u00ed como los principios de la funci\u00f3n administrativa y lo "
    "establecido en el Acuerdo 740 de 2019 (reglamentado en los art\u00edculos 147 a 160 del "
    "Decreto Distrital 642 de 2025); realiza el an\u00e1lisis t\u00e9cnico y jur\u00eddico de las "
    "siguientes solicitudes de modificaci\u00f3n contractual a contratos de prestaci\u00f3n "
    "de servicios profesionales y de apoyo a la gesti\u00f3n:"
)
CONCLUSION = (
    "A trav\u00e9s de este an\u00e1lisis, fortalecemos el acompa\u00f1amiento a la gesti\u00f3n local "
    "en la b\u00fasqueda de la mejora continua, siendo consecuentes en que cada "
    "modificaci\u00f3n propuesta responda a una necesidad real y sustentada del servicio, "
    "contribuyendo as\u00ed a una administraci\u00f3n m\u00e1s eficiente y responsable frente a "
    "los retos de la actual vigencia."
)


def solicitud(**cambios):
    fila = {
        "archivo": "solicitud-prueba.pdf",
        "localidad": "San Cristobal",
        "sipse": "990001",
        "contrato": "FDLSC-CPS-001-2026",
        "contratista": "CONTRATISTA DE PRUEBA",
        "supervisor": "SUPERVISOR DE PRUEBA",
        "objeto": "Prestar servicios profesionales de apoyo a la gestion local.",
        "fecha_inicio": "01/01/2026",
        "plazo_inicial": "9 meses",
        "fecha_terminacion_inicial": "30/09/2026",
        "prorroga_solicitada": "2 meses",
        "fecha_terminacion_final": "30/11/2026",
        "valor_inicial": 27000000,
        "valor_adicion": 6000000,
        "estado": "EN PROCESO DE ANALISIS",
    }
    fila.update(cambios)
    return fila


def generar(filas=None):
    return generar_documento_localidad(
        filas or [solicitud()], PLANTILLA, FECHA_RESPUESTA, "PROFESIONAL DE PRUEBA"
    )


def textos(doc):
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


class RespuestaWordRevision4Test(unittest.TestCase):
    def test_introduccion_y_cierre_revisados_en_todas_las_localidades(self):
        for localidad in ("San Cristobal", "La Candelaria", "Kennedy"):
            with self.subTest(localidad=localidad):
                parrafos = textos(Document(BytesIO(generar([solicitud(localidad=localidad)]))))
                self.assertEqual(parrafos.count(INTRODUCCION), 1)
                self.assertEqual(parrafos.count(CONCLUSION), 1)
                delegado = next(t for t in parrafos if t.startswith("El Fondo de Desarrollo Local tiene el deber"))
                self.assertIn("facultad delegada por el art\u00edculo 654 del Decreto 642 de 2025", delegado)
                self.assertLess(parrafos.index(delegado), parrafos.index(CONCLUSION))
                self.assertLess(parrafos.index(CONCLUSION), parrafos.index("Cordialmente,"))

    def test_supresiones_del_revisor_y_recomendaciones_no_modificadas(self):
        contenido = "\n".join(textos(Document(BytesIO(generar()))))
        for anterior in (
            "Decreto 411 de 2016", "Decreto Reglamentario 768 de 2019",
            "Decreto 168 de 2021", "Directiva Conjunta No. 01 de 2025",
            "Resoluci\u00f3n 467 de 2025", "emitir\u00e1 una \u00fanica respuesta",
        ):
            self.assertNotIn(anterior, contenido)
        self.assertIn("Radicado No. 20234500261013 Fecha: 21-07-2023", contenido)
        self.assertIn("El Fondo de Desarrollo Local debe tener la disponibilidad presupuestal", contenido)
        self.assertIn("Con una sustentada y soportada modificaci\u00f3n contractual", contenido)

    def test_nota_revisada_sin_listado_general(self):
        nota = next(t for t in textos(Document(BytesIO(generar()))) if t.startswith("Importante resaltar"))
        self.assertIn("radicado 20262100323333 del 31 de agosto de 2026, emiti\u00f3", nota)
        self.assertIn("ALERTAS TEMPRANAS SEGUNDO SEMESTRE VIGENCIA 2026", nota)
        self.assertIn("cumplimiento de t\u00e9rminos procedimentales", nota)
        self.assertTrue(nota.endswith("optimizar la gesti\u00f3n administrativa y mitigar riesgos institucionales."))
        self.assertNotIn("radicado c ", nota)
        self.assertNotIn("siendo su contenido es", nota)

    def test_nota_conserva_texto_revisado_al_actualizar_radicado_y_fecha(self):
        fila = solicitud(radicado_salida="20262100999999", fecha_salida="01/09/2026")
        nota = next(t for t in textos(Document(BytesIO(generar([fila])))) if t.startswith("Importante resaltar"))
        self.assertIn("radicado 20262100999999 del 1 de septiembre de 2026, emiti\u00f3", nota)
        self.assertNotIn("20262100323333", nota)
        self.assertIn("ALERTAS TEMPRANAS SEGUNDO SEMESTRE VIGENCIA 2026", nota)
        self.assertIn("herramientas preventivas y los lineamientos operativos", nota)

    def test_nota_con_fecha_de_respuesta_si_falta_fecha_salida(self):
        fila = solicitud(radicado_salida="20262100999999")
        nota = next(t for t in textos(Document(BytesIO(generar([fila])))) if t.startswith("Importante resaltar"))
        self.assertIn("radicado 20262100999999 del 17 de septiembre de 2026", nota)
        self.assertIn("SEGUNDO SEMESTRE VIGENCIA 2026", nota)

    def test_datos_dinamicos_y_trece_solicitudes_sin_copiar_ejemplos_del_word(self):
        filas = [solicitud(sipse=str(990001 + i), contrato=f"FDLSC-CPS-{i + 1:03}-2026") for i in range(13)]
        doc = Document(BytesIO(generar(filas)))
        parrafos = textos(doc)
        self.assertEqual(len(doc.tables), 13)
        self.assertEqual(
            [t for t in parrafos if t.startswith("SOLICITUD ")],
            [f"SOLICITUD {990001 + i}" for i in reversed(range(13))],
        )
        self.assertIn("Bogot\u00e1 D.C., 17 de septiembre de 2026", parrafos)
        self.assertTrue(any("Local de San Crist\u00f3bal" in t for t in parrafos))
        self.assertIn("Proyect\u00f3: PROFESIONAL DE PRUEBA - Profesional DGDL", parrafos)
        contenido = "\n".join(parrafos)
        for ejemplo in ("156648", "156596", "ANGELICA MARIA ANGARITA SERRANO", "La Candelaria"):
            self.assertNotIn(ejemplo, contenido)
        self.assertEqual(doc.tables[0].rows[2].cells[1].text, "FDLSC-CPS-013-2026")

    def test_formato_y_recursos_de_la_plantilla_se_conservan(self):
        resultado = generar()
        doc = Document(BytesIO(resultado))
        plantilla = Document(PLANTILLA)
        for referencia, actual in zip(plantilla.sections, doc.sections):
            for campo in ("page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin"):
                self.assertEqual(getattr(referencia, campo), getattr(actual, campo))
        for p in doc.paragraphs:
            if p.text.strip() in (INTRODUCCION, CONCLUSION) or p.text.startswith(("Importante resaltar", "El Fondo de Desarrollo Local tiene el deber")):
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
                self.assertTrue(all(r.font.name == "Garamond" for r in p.runs if r.text))
                self.assertTrue(all(r.font.size.pt == 11 for r in p.runs if r.text))
                self.assertTrue(all(not r.italic for r in p.runs if r.text))
        with ZipFile(PLANTILLA) as origen, ZipFile(BytesIO(resultado)) as salida:
            for nombre in origen.namelist():
                if nombre.startswith(("word/media/", "word/header", "word/footer")) and nombre.endswith((".xml", ".png", ".jpeg")):
                    esperado, actual = origen.read(nombre), salida.read(nombre)
                    if nombre.endswith(".xml"):
                        esperado = etree.tostring(etree.fromstring(esperado), method="c14n")
                        actual = etree.tostring(etree.fromstring(actual), method="c14n")
                    self.assertEqual(esperado, actual)
            self.assertFalse(any(n.startswith("word/comments") for n in salida.namelist()))
        for tag in ("ins", "del", "moveFrom", "moveTo", "commentRangeStart", "commentRangeEnd", "commentReference"):
            self.assertEqual(len(doc.element.findall(".//" + qn("w:" + tag))), 0)

    def test_zip_por_localidad_incluye_los_nuevos_textos(self):
        contenido, _, resumen = generar_zip_asistencia(
            [solicitud(), solicitud(localidad="La Candelaria", sipse="990002")],
            PLANTILLA, FECHA_RESPUESTA, "PROFESIONAL DE PRUEBA",
        )
        self.assertEqual(len(resumen), 2)
        with ZipFile(BytesIO(contenido)) as zip_final:
            archivos_word = [n for n in zip_final.namelist() if n.endswith(".docx")]
            self.assertEqual(len(archivos_word), 2)
            self.assertEqual(len([n for n in zip_final.namelist() if n.endswith(".xlsx")]), 1)
            for nombre in archivos_word:
                parrafos = textos(Document(BytesIO(zip_final.read(nombre))))
                self.assertIn(INTRODUCCION, parrafos)
                self.assertIn(CONCLUSION, parrafos)


if __name__ == "__main__":
    unittest.main()
