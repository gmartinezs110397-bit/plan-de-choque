from __future__ import annotations

import calendar
import copy
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable


LOCALIDADES_RADICADO = [
    "Usaquén",
    "Chapinero",
    "Santa Fe",
    "San Cristóbal",
    "Usme",
    "Tunjuelito",
    "Bosa",
    "Kennedy",
    "Fontibón",
    "Engativá",
    "Suba",
    "Barrios Unidos",
    "Teusaquillo",
    "Los Mártires",
    "Antonio Nariño",
    "Puente Aranda",
    "La Candelaria",
    "Rafael Uribe Uribe",
    "Ciudad Bolívar",
    "Sumapaz",
]

MESES_ES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)

PLAZO_SOLICITUD_WORD = "9 meses"
CIERRE_VIGENCIA_FISCAL_2026 = date(2026, 12, 31)

ESTADOS_ASISTENCIA = (
    "EN PROCESO DE ANÁLISIS",
    "EN PROCESO DE APROBACIÓN",
    "ORFEADO",
    "EN BANDEJA DEL FDL",
)


def _es_vacio(valor) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float) and math.isnan(valor):
        return True
    texto = str(valor).strip().lower()
    return texto in {"", "nan", "nat", "none"}


def _texto(valor) -> str:
    return "" if _es_vacio(valor) else str(valor)


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFD", _texto(texto).strip().lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto)


def _limpiar_texto(texto: str) -> str:
    if _es_vacio(texto):
        return ""
    texto = str(texto).replace("\xa0", " ")
    texto = re.sub(r"(?<=\w)-\s+(?=\w)", "-", texto)
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip(" \n\r\t:")


def _mes_numero(nombre: str) -> int | None:
    norm = _normalizar(nombre)
    for idx, mes in enumerate(MESES_ES, 1):
        if _normalizar(mes) == norm:
            return idx
    return None


def parsear_fecha(texto: str | date | datetime | None) -> date | None:
    if isinstance(texto, datetime):
        return texto.date()
    if isinstance(texto, date):
        return texto
    valor = _limpiar_texto(str(texto or ""))
    if not valor or valor.upper() in {"N/A", "NA"}:
        return None

    match = re.search(r"\b(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\b", valor)
    if match:
        dia, mes, anio = [int(x) for x in match.groups()]
        if anio < 100:
            anio += 2000
        try:
            return date(anio, mes, dia)
        except ValueError:
            return None

    match = re.search(
        r"\b(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})\b",
        valor,
        flags=re.IGNORECASE,
    )
    if match:
        dia = int(match.group(1))
        mes = _mes_numero(match.group(2))
        anio = int(match.group(3))
        if mes:
            try:
                return date(anio, mes, dia)
            except ValueError:
                return None
    return None


def formato_fecha_larga(fecha: date | datetime | None) -> str:
    parsed = parsear_fecha(fecha)
    if not parsed:
        return ""
    return f"{parsed.day} de {MESES_ES[parsed.month - 1]} de {parsed.year}"


def formato_fecha_corta(fecha: date | datetime | None) -> str:
    parsed = parsear_fecha(fecha)
    if not parsed:
        return ""
    return parsed.strftime("%d/%m/%Y")


def formato_moneda(valor) -> str:
    numero = _a_numero(valor)
    if numero is None:
        return ""
    return "$" + f"{int(round(numero)):,.0f}".replace(",", ".")


def _a_numero(valor) -> int | None:
    if _es_vacio(valor):
        return None
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and math.isnan(valor):
            return None
        return int(round(float(valor)))
    texto = str(valor)
    if not texto.strip():
        return None
    texto = texto.replace("$", "").replace(".", "").replace(",", "")
    texto = re.sub(r"[^\d-]", "", texto)
    if not texto:
        return None
    try:
        return int(texto)
    except ValueError:
        return None


def _extraer_entre(texto: str, etiqueta: str, siguientes: Iterable[str]) -> str:
    cierre = "|".join(re.escape(s) for s in siguientes)
    patron = rf"{re.escape(etiqueta)}\s*:?\s*(.*?)(?=\s+(?:{cierre})(?:[^:]{{0,80}})?:|\s+[IVX]+\.\s|$)"
    match = re.search(patron, texto, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return ""
    return _limpiar_texto(match.group(1))


def _extraer_seccion(texto: str, inicio_patron: str, fin_patrones: Iterable[str]) -> str:
    inicio = re.search(inicio_patron, texto, flags=re.IGNORECASE | re.DOTALL)
    if not inicio:
        return texto
    desde = inicio.end()
    fin = len(texto)
    for patron in fin_patrones:
        match = re.search(patron, texto[desde:], flags=re.IGNORECASE | re.DOTALL)
        if match:
            fin = min(fin, desde + match.start())
    return texto[desde:fin]


def _extraer_fecha_entre(texto: str, etiqueta: str, siguientes: Iterable[str]) -> str:
    return formato_fecha_corta(parsear_fecha(_extraer_entre(texto, etiqueta, siguientes)))


def _limpiar_objeto_solicitud(texto: str) -> str:
    objeto = _limpiar_texto(texto).strip(' "“”')
    if not objeto:
        return ""

    cortes = (
        r"\bContratista\b",
        r"\bSupervisor\b",
        r"\bValor\s+Inicial\b",
        r"\bN[uú]mero\s+del\s+proceso\b",
        r"\bFecha\s+de\s+publicaci[oó]n\b",
        r"\bLink\s+del\s+proceso\b",
        r"\bSolicitud\s+de\s+modificaci[oó]n\s+contractual\b",
        r"\bEstado\s+financiero\s+del\s+contrato\b",
        r"\bValor\s+del\s+contrato\b",
        r"\bValor\s+ejecutado\b",
        r"\bValor\s+pagado\b",
        r"\bValor\s+a\s+adicionar\b",
        r"\bC[oó]digo\s+postal\b",
        r"\bInformaci[oó]n\s+l[ií]nea\b",
        r"\bwww\.",
        r"\bC[oó]digo\s*:",
        r"\bVersi[oó]n\s*:",
        r"\bVigencia\s*:",
        r"\bCaso\s+HOLA\b",
        r"\bP[aá]gina\b",
    )
    posiciones = []
    for patron in cortes:
        match = re.search(patron, objeto, flags=re.IGNORECASE)
        if match:
            posiciones.append(match.start())
    if posiciones:
        objeto = objeto[: min(posiciones)]
    return _limpiar_texto(objeto).strip(' "“”')


def _extraer_objeto_solicitud(seccion_resumen: str) -> str:
    objeto = _extraer_entre(
        seccion_resumen,
        "Objeto",
        [
            "Contratista",
            "Supervisor",
            "Valor Inicial",
            "Número del proceso",
            "Numero del proceso",
            "Fecha de publicación",
            "Fecha de publicacion",
            "Link del proceso",
        ],
    )
    if objeto:
        return _limpiar_objeto_solicitud(objeto)

    cierre = "|".join(
        re.escape(s)
        for s in (
            "Contratista",
            "Supervisor",
            "Valor Inicial",
            "Número del proceso",
            "Numero del proceso",
            "Fecha de publicación",
            "Fecha de publicacion",
            "Link del proceso",
        )
    )
    match = re.search(
        rf"\bObjeto\s*:?\s*(.*?)(?=\s+(?:{cierre})\b|\s+[IVX]+\.\s|$)",
        seccion_resumen,
        flags=re.IGNORECASE | re.DOTALL,
    )
    return _limpiar_objeto_solicitud(match.group(1)) if match else ""


def _duracion_desde_texto(texto: str) -> tuple[int, int]:
    texto_norm = _normalizar(texto)
    meses = 0
    dias = 0
    match = re.search(r"\((\d+)\)\s*mes", texto_norm)
    if match:
        meses = int(match.group(1))
    else:
        match = re.search(r"(\d+)\s*mes", texto_norm)
        if match:
            meses = int(match.group(1))
    match = re.search(r"(\d+)\s*d[ií]a", texto_norm)
    if match:
        dias = int(match.group(1))
    return meses, dias


def _formato_duracion(meses: int, dias: int) -> str:
    partes: list[str] = []
    if meses:
        partes.append(f"{meses} mes" + ("es" if meses != 1 else ""))
    if dias:
        partes.append(f"{dias} día" + ("s" if dias != 1 else ""))
    return " y ".join(partes) if partes else ""


def _datos_calculadora(fila: dict) -> dict:
    plazo_meses, plazo_dias = _duracion_desde_texto(_texto(fila.get("plazo_inicial")))
    prorroga_meses, prorroga_dias = _duracion_desde_texto(_texto(fila.get("prorroga_solicitada")))
    return {
        "fecha_inicio": parsear_fecha(fila.get("fecha_inicio")) or _texto(fila.get("fecha_inicio")),
        "plazo_meses": plazo_meses,
        "plazo_dias": plazo_dias,
        "valor_total": _a_numero(fila.get("valor_inicial_texto")) or _a_numero(fila.get("valor_inicial")),
        "prorroga_meses": prorroga_meses,
        "prorroga_dias": prorroga_dias,
        "texto_plazo": _texto(fila.get("plazo_inicial")),
        "texto_prorroga": _texto(fila.get("prorroga_solicitada")),
        "texto_valor": _texto(fila.get("valor_inicial_texto")) or formato_moneda(fila.get("valor_inicial")),
    }


def _sumar_meses(fecha: date, meses: int) -> date:
    mes_base = fecha.month - 1 + meses
    anio = fecha.year + mes_base // 12
    mes = mes_base % 12 + 1
    dia = min(fecha.day, calendar.monthrange(anio, mes)[1])
    return date(anio, mes, dia)


def calcular_fecha_fin_inicial(fecha_inicio: date | None, plazo_texto: str) -> date | None:
    if not fecha_inicio:
        return None
    meses, dias = _duracion_desde_texto(plazo_texto)
    if meses == 0 and dias == 0:
        return None
    fin = _sumar_meses(fecha_inicio, meses)
    return date.fromordinal(fin.toordinal() - 1 + dias)


def calcular_fecha_fin_prorroga(fecha_fin_inicial: date | None, prorroga_texto: str) -> date | None:
    if not fecha_fin_inicial:
        return None
    meses, dias = _duracion_desde_texto(prorroga_texto)
    if meses == 0 and dias == 0:
        return None
    fin = _sumar_meses(fecha_fin_inicial, meses)
    return date.fromordinal(fin.toordinal() + dias)


def _ultimo_dia_mes(fecha: date) -> date:
    return date(fecha.year, fecha.month, calendar.monthrange(fecha.year, fecha.month)[1])


def _primer_dia_mes_siguiente(fecha: date) -> date:
    if fecha.month == 12:
        return date(fecha.year + 1, 1, 1)
    return date(fecha.year, fecha.month + 1, 1)


def _ultimo_dia_mes_anterior(fecha: date) -> date:
    mes = fecha.month - 1
    anio = fecha.year
    if mes == 0:
        mes = 12
        anio -= 1
    return _ultimo_dia_mes(date(anio, mes, 1))


def _meses_completos(inicio: date, fin: date) -> int:
    return (fin.year * 12 + fin.month) - (inicio.year * 12 + inicio.month) + 1


def _calcular_validacion_solicitud(fila: dict) -> dict:
    fecha_inicio = parsear_fecha(fila.get("fecha_inicio"))
    fecha_fin_solicitud = parsear_fecha(fila.get("fecha_terminacion_inicial"))
    fecha_final_solicitud = parsear_fecha(fila.get("fecha_terminacion_final"))
    plazo_texto = _texto(fila.get("plazo_inicial")) or PLAZO_SOLICITUD_WORD
    prorroga_texto = _texto(fila.get("prorroga_solicitada"))
    plazo_meses, plazo_dias = _duracion_desde_texto(plazo_texto)
    prorroga_meses, prorroga_dias = _duracion_desde_texto(prorroga_texto)
    valor_inicial = _a_numero(fila.get("valor_inicial_texto")) or _a_numero(fila.get("valor_inicial"))
    valor_adicion = _a_numero(fila.get("valor_adicion_texto")) or _a_numero(fila.get("valor_adicion"))

    fecha_fin_calculada = calcular_fecha_fin_inicial(fecha_inicio, plazo_texto)
    fecha_final_calculada = calcular_fecha_fin_prorroga(fecha_fin_calculada, prorroga_texto)
    prorroga_dias_ajustada = prorroga_dias
    fecha_final_ajustada = fecha_final_calculada
    ajuste_dia_31 = False
    if fecha_final_calculada and fecha_final_calculada.day == 30:
        if calendar.monthrange(fecha_final_calculada.year, fecha_final_calculada.month)[1] == 31:
            ajuste_dia_31 = True
            prorroga_dias_ajustada += 1
            fecha_final_ajustada = date.fromordinal(fecha_final_calculada.toordinal() + 1)

    valor_adicion_teorico = None
    if (
        fecha_fin_calculada
        and fecha_final_ajustada
        and valor_inicial is not None
        and (plazo_meses or plazo_dias)
    ):
        dias_plazo = 30 * plazo_meses + plazo_dias
        if dias_plazo:
            tarifa_dia = valor_inicial / dias_plazo
            tarifa_mes = tarifa_dia * 30
            inicio_prorroga = date.fromordinal(fecha_fin_calculada.toordinal() + 1)
            primer_mes_completo = (
                inicio_prorroga
                if inicio_prorroga.day == 1
                else _primer_dia_mes_siguiente(inicio_prorroga)
            )
            ultimo_mes_completo = (
                fecha_final_ajustada
                if fecha_final_ajustada == _ultimo_dia_mes(fecha_final_ajustada)
                else _ultimo_dia_mes_anterior(fecha_final_ajustada)
            )

            meses_completos = (
                _meses_completos(primer_mes_completo, ultimo_mes_completo)
                if primer_mes_completo <= ultimo_mes_completo
                else 0
            )
            dias_antes = (
                0
                if inicio_prorroga.day == 1
                else _ultimo_dia_mes(inicio_prorroga).toordinal() - inicio_prorroga.toordinal() + 1
            )
            dias_despues = (
                0
                if fecha_final_ajustada == _ultimo_dia_mes(fecha_final_ajustada)
                else fecha_final_ajustada.day
            )
            if meses_completos == 0:
                if inicio_prorroga.month == fecha_final_ajustada.month and inicio_prorroga.year == fecha_final_ajustada.year:
                    ajuste_proporcional = 1 if calendar.monthrange(fecha_final_ajustada.year, fecha_final_ajustada.month)[1] == 31 else 0
                else:
                    ajuste_proporcional = (
                        (1 if calendar.monthrange(inicio_prorroga.year, inicio_prorroga.month)[1] == 31 else 0)
                        + (1 if fecha_final_ajustada.day == 31 else 0)
                    )
                dias_proporcionales = (
                    fecha_final_ajustada.toordinal() - inicio_prorroga.toordinal() + 1 - ajuste_proporcional
                )
            else:
                ajuste_proporcional = (
                    (1 if dias_antes > 0 and calendar.monthrange(inicio_prorroga.year, inicio_prorroga.month)[1] == 31 else 0)
                    + (1 if dias_despues > 0 and fecha_final_ajustada.day == 31 else 0)
                )
                dias_proporcionales = dias_antes + dias_despues - ajuste_proporcional
            valor_adicion_teorico = round(meses_completos * tarifa_mes + dias_proporcionales * tarifa_dia)

    prorroga_ajustada_texto = _formato_duracion(prorroga_meses, prorroga_dias_ajustada)
    return {
        "fecha_fin_calculada": fecha_fin_calculada,
        "fecha_final_calculada": fecha_final_calculada,
        "fecha_final_ajustada": fecha_final_ajustada,
        "prorroga_ajustada_texto": prorroga_ajustada_texto,
        "valor_adicion_teorico": valor_adicion_teorico,
        "se_ajusta_fecha_fin": bool(fecha_fin_solicitud and fecha_fin_calculada and fecha_fin_solicitud == fecha_fin_calculada),
        "se_ajusta_prorroga": not ajuste_dia_31 and bool(prorroga_texto.strip()),
        "se_ajusta_fecha_final": bool(fecha_final_solicitud and fecha_final_calculada and fecha_final_solicitud == fecha_final_calculada),
        "se_ajusta_adicion": bool(
            valor_adicion is not None
            and valor_adicion_teorico is not None
            and abs(valor_adicion - valor_adicion_teorico) <= 1
        ),
    }


def _canon_localidad(texto: str) -> str:
    norm = _normalizar(texto)
    norm = re.sub(r"^(la|los|las|el)\s+", "", norm)
    alias = {
        "usaquen": "Usaquén",
        "chapinero": "Chapinero",
        "santa fe": "Santa Fe",
        "san cristobal": "San Cristóbal",
        "usme": "Usme",
        "tunjuelito": "Tunjuelito",
        "bosa": "Bosa",
        "kennedy": "Kennedy",
        "fontibon": "Fontibón",
        "engativa": "Engativá",
        "suba": "Suba",
        "barrios unidos": "Barrios Unidos",
        "teusaquillo": "Teusaquillo",
        "martires": "Los Mártires",
        "antonio narino": "Antonio Nariño",
        "puente aranda": "Puente Aranda",
        "candelaria": "La Candelaria",
        "rafael uribe": "Rafael Uribe Uribe",
        "rafael uribe uribe": "Rafael Uribe Uribe",
        "ciudad bolivar": "Ciudad Bolívar",
        "sumapaz": "Sumapaz",
    }
    return alias.get(norm, texto.strip())


def _extraer_localidad(texto: str) -> str:
    match = re.search(
        r"Alcald[ií]a\s+Local\s+de\s+(?:la\s+|los\s+|las\s+|el\s+)?([A-Za-zÁÉÍÓÚÜÑáéíóúüñ\s]+)",
        texto,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    candidato = _limpiar_texto(match.group(1).split("Edificio")[0])
    candidato = candidato.split("Código")[0].strip()
    return _canon_localidad(candidato)


def _numero_contrato_corto(numero: str) -> str:
    texto = _limpiar_texto(numero).upper()
    match = re.search(r"(?:CPS|CPS-P|CD|SAMC|FDL[A-Z-]*)[-\s]*(\d{1,4})[-/](20\d{2})", texto)
    if match:
        return f"{match.group(1).zfill(3)}-{match.group(2)}"
    match = re.search(r"(\d{1,4})[-/](20\d{2})", texto)
    if match:
        return f"{match.group(1).zfill(3)}-{match.group(2)}"
    return texto


def _extraer_sipse(nombre_archivo: str, texto: str) -> str:
    candidatos = re.findall(r"\b1\d{5}\b", nombre_archivo)
    if candidatos:
        return candidatos[0]
    candidatos = re.findall(r"SIPSE\s*:?\s*(1\d{5})", texto, flags=re.IGNORECASE)
    return candidatos[0] if candidatos else ""


def extraer_texto_pdf(nombre_archivo: str, contenido: bytes) -> tuple[str, list[str]]:
    errores: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(contenido))
        paginas = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(paginas), errores
    except Exception as exc:  # pragma: no cover - depende del PDF recibido.
        errores.append(f"{nombre_archivo}: no se pudo leer el PDF ({exc}).")
        return "", errores


def analizar_solicitud_pdf(nombre_archivo: str, contenido: bytes) -> tuple[dict, list[str]]:
    texto_pdf, errores = extraer_texto_pdf(nombre_archivo, contenido)
    texto = re.sub(r"\s+", " ", texto_pdf.replace("\xa0", " "))
    seccion_resumen = _extraer_seccion(
        texto,
        r"\bI\.\s*RESUMEN\s+CONTRACTUAL\b",
        [r"\bII\.\s*INFORMACI[oó]N\s+DE\s+LA\s+MODIFICACI[oó]N\s+SOLICITADA\b"],
    )
    seccion_modificacion = _extraer_seccion(
        texto,
        r"\bII\.\s*INFORMACI[oó]N\s+DE\s+LA\s+MODIFICACI[oó]N\s+SOLICITADA\b",
        [
            r"\bIII\.\s*INFORMACI[oó]N\s+DE\s+MODIFICACIONES\s+ANTERIORES\b",
            r"\bESTADO\s+FINANCIERO\b",
            r"\bSOLICITUD\s+DE\s+MODIFICACI[oó]N\s+CONTRACTUAL\b",
        ],
    )
    fila: dict = {
        "archivo": nombre_archivo,
        "sipse": _extraer_sipse(nombre_archivo, texto),
        "localidad": _extraer_localidad(texto),
        "fecha_solicitud": _extraer_fecha_entre(texto, "Fecha de Solicitud", ["Área de Origen", "Area de Origen"]),
        "area_origen": _extraer_entre(texto, "Área de Origen", ["I. RESUMEN CONTRACTUAL", "Número Contrato"]),
        "contrato": _extraer_entre(seccion_resumen, "Número Contrato", ["Fecha de Suscripción"]),
        "fecha_suscripcion": _extraer_entre(seccion_resumen, "Fecha de Suscripción", ["Tipo de Contrato"]),
        "tipo_contrato": _extraer_entre(seccion_resumen, "Tipo de Contrato", ["Plazo Inicial"]),
        "plazo_inicial": _extraer_entre(seccion_resumen, "Plazo Inicial", ["Fecha de Inicio"]),
        "fecha_inicio": _extraer_entre(seccion_resumen, "Fecha de Inicio", ["Fecha de Terminación Inicial"]),
        "fecha_terminacion_inicial": _extraer_entre(seccion_resumen, "Fecha de Terminación Inicial", ["Objeto"]),
        "objeto": _extraer_objeto_solicitud(seccion_resumen),
        "contratista": _extraer_entre(seccion_resumen, "Contratista", ["Supervisor"]),
        "supervisor": _extraer_entre(seccion_resumen, "Supervisor", ["Valor Inicial"]),
        "valor_inicial_texto": _extraer_entre(seccion_resumen, "Valor Inicial", ["Número del proceso", "Numero del proceso"]),
        "valor_inicial": _a_numero(_extraer_entre(seccion_resumen, "Valor Inicial", ["Número del proceso", "Numero del proceso"])),
        "proceso_secop": _extraer_entre(seccion_resumen, "Número del proceso SECOP I o II", ["Fecha de publicación"]),
        "fecha_publicacion_secop": _extraer_entre(seccion_resumen, "Fecha de publicación del proceso en SECOP I o II", ["Link del proceso"]),
        "link_secop": "",
        "valor_adicion_texto": "",
        "valor_adicion": None,
        "prorroga_solicitada": "",
        "fecha_terminacion_final": "",
        "estado": ESTADOS_ASISTENCIA[0],
        "observaciones": "NINGUNA",
    }

    url = re.search(r"https?://\S+", texto_pdf)
    if url:
        fila["link_secop"] = url.group(0).strip()

    valor_adicion = re.search(r"Valor\s+a\s+Adicionar\s*:\s*(\$?\s*[\d\.\,]+)", seccion_modificacion, flags=re.IGNORECASE)
    if not valor_adicion:
        valor_adicion = re.search(
            r"Adici[oó]n\s+Valor\s+Pr[oó]rroga\s+Tiempo.*?(\$\s*[\d\.\,]+)",
            seccion_modificacion,
            flags=re.IGNORECASE | re.DOTALL,
        )
    if valor_adicion:
        fila["valor_adicion_texto"] = _limpiar_texto(valor_adicion.group(1))
        fila["valor_adicion"] = _a_numero(valor_adicion.group(1))

    match_prorroga = re.search(
        r"Tiempo\s*:\s*(.*?)(?:Adici[oó]n\s+y\s+Pr[oó]rroga|Fecha\s+Terminaci[oó]n\s+Final|III\.)",
        seccion_modificacion,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if match_prorroga:
        fila["prorroga_solicitada"] = _limpiar_texto(match_prorroga.group(1))

    match_fecha_final = re.search(
        r"Fecha\s+Terminaci[oó]n\s+Final\s*:?\s*(.*?)\s*$",
        seccion_modificacion,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if match_fecha_final:
        fila["fecha_terminacion_final"] = _limpiar_texto(match_fecha_final.group(1))

    for campo in ("contrato", "objeto", "contratista", "supervisor", "plazo_inicial", "prorroga_solicitada"):
        fila[campo] = _limpiar_texto(fila.get(campo, ""))
    fila["contrato"] = fila["contrato"].replace(" ", "")
    fila["objeto"] = fila["objeto"].strip(' "“”')

    requeridos = ("localidad", "contrato", "contratista", "fecha_inicio", "plazo_inicial")
    faltantes = [campo for campo in requeridos if not fila.get(campo)]
    if faltantes:
        errores.append(f"{nombre_archivo}: revise {', '.join(faltantes)}.")
    return fila, errores


def analizar_solicitudes(archivos: Iterable[tuple[str, bytes]]) -> tuple[list[dict], list[str]]:
    filas: list[dict] = []
    errores: list[str] = []
    for nombre, contenido in archivos:
        fila, err = analizar_solicitud_pdf(nombre, contenido)
        filas.append(fila)
        errores.extend(err)
    return filas, errores


def extraer_radicados(contenido_pdf: bytes) -> tuple[dict[str, str], date | None, list[str]]:
    texto, errores = extraer_texto_pdf("Listado de radicado", contenido_pdf)
    numeros = re.findall(r"\b20\d{12}\b", texto)
    mapa = {
        localidad: numeros[idx]
        for idx, localidad in enumerate(LOCALIDADES_RADICADO)
        if idx < len(numeros)
    }

    fecha = None
    match = re.search(r"Fecha\s*:\s*(\d{1,2})-(\d{1,2})-(\d{1,4})", texto, flags=re.IGNORECASE)
    if match:
        dia = int(match.group(1))
        mes = int(match.group(2))
        anio_txt = match.group(3)
        anio = int(anio_txt) if len(anio_txt) == 4 else None
        if not anio and numeros:
            anio = int(numeros[0][:4])
        if anio:
            try:
                fecha = date(anio, mes, dia)
            except ValueError:
                fecha = None
    if not mapa:
        errores.append("No se encontraron radicados de salida en el archivo general.")
    return mapa, fecha, errores


def enriquecer_con_radicados(filas: list[dict], mapa: dict[str, str], fecha_salida: date | None) -> list[dict]:
    resultado: list[dict] = []
    for fila in filas:
        copia = dict(fila)
        localidad = _canon_localidad(copia.get("localidad", ""))
        copia["localidad"] = localidad
        copia["radicado_salida"] = _texto(copia.get("radicado_salida")) or mapa.get(localidad, "")
        copia["fecha_salida"] = _texto(copia.get("fecha_salida")) or formato_fecha_corta(fecha_salida)
        resultado.append(copia)
    return resultado


def _orden_localidad(localidad: str) -> int:
    canon = _canon_localidad(localidad)
    try:
        return LOCALIDADES_RADICADO.index(canon)
    except ValueError:
        return 99


def _clave_sipse_desc(fila: dict) -> tuple[int, int, str]:
    sipse = _texto(fila.get("sipse")).strip()
    if sipse.isdigit():
        return (0, -int(sipse), "")
    return (1, 0, sipse)


def ordenar_filas(filas: Iterable[dict]) -> list[dict]:
    return sorted(
        (dict(f) for f in filas),
        key=lambda f: (
            _orden_localidad(f.get("localidad", "")),
            _clave_sipse_desc(f),
            _texto(f.get("contrato")),
        ),
    )


def _nombre_archivo_seguro(texto: str) -> str:
    texto = _normalizar(texto)
    texto = re.sub(r"[^a-z0-9]+", "_", texto).strip("_")
    return texto or "archivo"


def _descripcion_solicitud(fila: dict) -> str:
    tiene_adicion = (_a_numero(fila.get("valor_adicion")) or 0) > 0
    tiene_prorroga = bool(_texto(fila.get("prorroga_solicitada")).strip())
    if tiene_adicion and tiene_prorroga:
        return "ADICIÓN Y PRÓRROGA"
    if tiene_adicion:
        return "ADICIÓN"
    if tiene_prorroga:
        return "PRÓRROGA"
    return "MODIFICACIÓN"


def _mes_reporte(fila: dict) -> str:
    fecha = parsear_fecha(fila.get("fecha_solicitud"))
    return MESES_ES[fecha.month - 1].upper() if fecha else ""


def _nombre_hoja_contrato(fila: dict, consecutivo: int, usados: set[str]) -> str:
    contrato = _numero_contrato_corto(_texto(fila.get("contrato")))
    base = re.sub(r"[\[\]\*\?/\\:]", " ", contrato or f"Contrato {consecutivo}")
    base = re.sub(r"\s+", " ", base).strip()[:25] or f"Contrato {consecutivo}"
    nombre = base
    sufijo = 2
    while nombre in usados:
        extra = f" {sufijo}"
        nombre = f"{base[:31 - len(extra)]}{extra}"
        sufijo += 1
    usados.add(nombre)
    return nombre


def _agregar_ficha_contrato(wb, fila: dict, consecutivo: int, profesional: str) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    nombre = _nombre_hoja_contrato(fila, consecutivo, set(wb.sheetnames))
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False

    rojo = PatternFill("solid", fgColor="C00000")
    amarillo = PatternFill("solid", fgColor="FFC000")
    gris = PatternFill("solid", fgColor="D9D9D9")
    gris = PatternFill("solid", fgColor="D9D9D9")
    gris_claro = PatternFill("solid", fgColor="F3F4F6")
    borde = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )

    ws.merge_cells("A1:D1")
    ws["A1"] = f"FICHA DE REVISIÓN SECOP - {_numero_contrato_corto(_texto(fila.get('contrato')))}"
    ws["A1"].fill = rojo
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    secciones = [
        (
            "Información de la solicitud",
            [
                ("Localidad", _canon_localidad(fila.get("localidad", "")).upper()),
                ("SIPSE", _texto(fila.get("sipse"))),
                ("Radicado de salida", _texto(fila.get("radicado_salida"))),
                ("Fecha salida DGDL", parsear_fecha(fila.get("fecha_salida"))),
                ("Archivo origen", _texto(fila.get("archivo"))),
                ("Profesional DGDL", _texto(profesional)),
            ],
        ),
        (
            "Información del contrato",
            [
                ("No. de contrato", _texto(fila.get("contrato"))),
                ("Contratista", _texto(fila.get("contratista"))),
                ("Supervisor(a)", _texto(fila.get("supervisor"))),
                ("Objeto", _texto(fila.get("objeto"))),
                ("Fecha de inicio", _texto(fila.get("fecha_inicio"))),
                ("Plazo inicial", _texto(fila.get("plazo_inicial"))),
                ("Valor inicial", _texto(fila.get("valor_inicial_texto")) or formato_moneda(fila.get("valor_inicial"))),
                ("Fecha terminación inicial", _texto(fila.get("fecha_terminacion_inicial"))),
            ],
        ),
        (
            "Modificación solicitada",
            [
                ("Tipo de solicitud", _descripcion_solicitud(fila)),
                ("Prórroga solicitada", _texto(fila.get("prorroga_solicitada"))),
                ("Valor adición", _texto(fila.get("valor_adicion_texto")) or formato_moneda(fila.get("valor_adicion"))),
                ("Fecha terminación con prórroga", _texto(fila.get("fecha_terminacion_final"))),
                ("Estado / trámite", (_texto(fila.get("estado")) or ESTADOS_ASISTENCIA[0]).upper()),
                ("Observaciones", _texto(fila.get("observaciones")) or "NINGUNA"),
            ],
        ),
        (
            "Verificación SECOP por Kate",
            [
                ("Se ajusta en SECOP", ""),
                ("Observación SECOP", ""),
                ("Estado del cargue SECOP", ""),
                ("Fecha de revisión", ""),
                ("Proceso SECOP", _texto(fila.get("proceso_secop"))),
                ("Fecha publicación SECOP", parsear_fecha(fila.get("fecha_publicacion_secop"))),
                ("Link SECOP", _texto(fila.get("link_secop"))),
            ],
        ),
    ]

    fila_excel = 3
    for titulo, campos in secciones:
        ws.merge_cells(start_row=fila_excel, start_column=1, end_row=fila_excel, end_column=4)
        celda_titulo = ws.cell(fila_excel, 1, titulo)
        celda_titulo.fill = amarillo
        celda_titulo.font = Font(bold=True)
        celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
        celda_titulo.border = borde
        for col in range(1, 5):
            ws.cell(fila_excel, col).border = borde
        fila_excel += 1

        for etiqueta, valor in campos:
            ws.cell(fila_excel, 1, etiqueta)
            ws.cell(fila_excel, 2, valor)
            ws.merge_cells(start_row=fila_excel, start_column=2, end_row=fila_excel, end_column=4)
            for col in range(1, 5):
                celda = ws.cell(fila_excel, col)
                celda.border = borde
                celda.alignment = Alignment(vertical="center", wrap_text=True)
                if col == 1:
                    celda.fill = gris
                    celda.font = Font(bold=True)
                elif titulo == "Verificación SECOP por Kate" and etiqueta in {
                    "Se ajusta en SECOP",
                    "Observación SECOP",
                    "Estado del cargue SECOP",
                    "Fecha de revisión",
                }:
                    celda.fill = gris_claro
                if isinstance(celda.value, date):
                    celda.number_format = "DD/MM/YYYY"
                if etiqueta in {"Valor inicial", "Valor adición"}:
                    celda.number_format = '"$"#,##0'
            fila_excel += 1
        fila_excel += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28
    for row in range(1, fila_excel):
        ws.row_dimensions[row].height = 24
    for row in range(1, fila_excel):
        if ws.cell(row, 1).value in {"Objeto", "Observaciones", "Link SECOP", "Proceso SECOP"}:
            ws.row_dimensions[row].height = 72
    ws.freeze_panes = "A3"


def generar_excel_asistencia(filas: Iterable[dict], profesional: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws_aux = wb.active
    ws_aux.title = "Hoja1"
    for idx, estado in enumerate(ESTADOS_ASISTENCIA, 5):
        ws_aux.cell(idx, 3, estado)
    ws_listas = wb.create_sheet("Hoja2")
    for idx, localidad in enumerate(LOCALIDADES_RADICADO, 3):
        ws_listas.cell(idx, 3, _normalizar(localidad).upper())
    ws = wb.create_sheet("CPS 2026")
    wb.active = wb.sheetnames.index("CPS 2026")
    ws.sheet_view.showGridLines = False

    rojo = PatternFill("solid", fgColor="C00000")
    amarillo = PatternFill("solid", fgColor="FFC000")
    gris = PatternFill("solid", fgColor="D9D9D9")
    borde = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )

    ws.merge_cells("C1:R1")
    ws["C1"] = "MATRIZ DE SEGUIMIENTO ASISTENCIA TÉCNICA DE LOS FDL - ADICIONES Y/O PRÓRROGA"
    ws["C1"].fill = rojo
    ws["C1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("S1:T6")
    ws["S1"] = "Código: GET-AGL-F009\nVersión: 01\nVigencia:  27 de septiembre de 2022\nCaso HOLA: 268037"
    ws["S1"].alignment = Alignment(wrap_text=True, vertical="top")

    encabezados = [
        "CONSECUTIVO No.",
        "LOCALIDAD",
        "NÚMERO DE ORDEN",
        "FECHA DE INGRESO\nDD-MM-AAAA",
        "RADICADO DE ENTRADA FDL - SIPSE",
        "DEVUELTO SI / NO",
        "FECHA DE DEVOLUCIÓN\nDD-MM-AAAA",
        "DESCRIPCIÓN DE LA SOLICITUD ADICIÓN Y/O PRÓRROGA",
        "No. DE CONTRATO (EJ: XXX/2026)",
        "OBJETO ",
        "NOMBRES Y APELLIDOS CONTRATISTA CPS",
        "PLAZO INICIAL (MESES)",
        "VALOR INICIAL DEL CONTRATO ",
        "PLAZO DE LA PRÓRROGA SOLICITADA (MESES) ",
        "VALOR ADICIÓN SOLICITADA",
        "ESTADO / TRÁMITE",
        "RADICADO DE SALIDA",
        "FECHA SALIDA DGDL\nDD-MM-AAAA",
        "OBSERVACIONES",
        "MES A MES",
        "",
        "VENCE",
        "PROFESIONAL RESPONSABLE",
        "FECHA DIRECTORA",
        "DIAS HABILES",
        "DIAS SIN RESPUESTA",
    ]
    for col, titulo in enumerate(encabezados, 1):
        celda = ws.cell(7, col, titulo)
        celda.fill = amarillo
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde

    filas_ordenadas = ordenar_filas(filas)
    orden_por_localidad: dict[str, int] = {}
    for idx, fila in enumerate(filas_ordenadas, 8):
        localidad = _canon_localidad(fila.get("localidad", ""))
        orden_por_localidad[localidad] = orden_por_localidad.get(localidad, 0) + 1
        valores = [
            idx - 7,
            localidad.upper(),
            orden_por_localidad[localidad],
            parsear_fecha(fila.get("fecha_solicitud")),
            _texto(fila.get("sipse")),
            "NO",
            "N/A",
            _descripcion_solicitud(fila),
            _numero_contrato_corto(_texto(fila.get("contrato"))),
            _texto(fila.get("objeto")),
            _texto(fila.get("contratista")),
            _texto(fila.get("plazo_inicial")),
            _texto(fila.get("valor_inicial_texto")) or formato_moneda(fila.get("valor_inicial")),
            _texto(fila.get("prorroga_solicitada")),
            _texto(fila.get("valor_adicion_texto")) or formato_moneda(fila.get("valor_adicion")),
            (_texto(fila.get("estado")) or ESTADOS_ASISTENCIA[0]).upper(),
            "",
            "",
            (_texto(fila.get("observaciones")) or "NINGUNA").upper(),
            "",
            "",
            "",
            str(profesional or "").upper(),
            "",
            "",
            "",
        ]
        for col, valor in enumerate(valores, 1):
            celda = ws.cell(idx, col, valor)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=col in {10, 11})
            if col in {4, 18} and isinstance(valor, date):
                celda.number_format = "DD/MM/YYYY"
            if col in {13, 15}:
                celda.number_format = '"$"#,##0'

    ws_calc = wb.create_sheet("Calculadora")
    ws_calc.sheet_view.showGridLines = False
    ws_calc.merge_cells("A1:L1")
    ws_calc["A1"] = "DATOS PARA VERIFICAR EN CALCULADORA"
    ws_calc["A1"].fill = rojo
    ws_calc["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws_calc["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_calc.merge_cells("A2:L2")
    ws_calc["A2"] = (
        "Entradas tomadas de cada solicitud. Revise estos datos contra el PDF antes de validar en SECOP."
    )
    ws_calc["A2"].alignment = Alignment(wrap_text=True, vertical="center")

    encabezados_calc = [
        "No.",
        "Localidad",
        "Contrato",
        "Fecha de inicio",
        "Plazo - Meses",
        "Plazo - Días",
        "Valor total del contrato (COP)",
        "Prórroga - Meses",
        "Prórroga - Días",
        "Texto plazo en solicitud",
        "Texto prórroga en solicitud",
        "Valor inicial en solicitud",
    ]
    for col, titulo in enumerate(encabezados_calc, 1):
        celda = ws_calc.cell(4, col, titulo)
        celda.fill = amarillo
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde

    for idx, fila in enumerate(filas_ordenadas, 5):
        localidad = _canon_localidad(fila.get("localidad", ""))
        datos_calc = _datos_calculadora(fila)
        valores_calc = [
            idx - 4,
            localidad.upper(),
            _numero_contrato_corto(_texto(fila.get("contrato"))),
            datos_calc["fecha_inicio"],
            datos_calc["plazo_meses"],
            datos_calc["plazo_dias"],
            datos_calc["valor_total"],
            datos_calc["prorroga_meses"],
            datos_calc["prorroga_dias"],
            datos_calc["texto_plazo"],
            datos_calc["texto_prorroga"],
            datos_calc["texto_valor"],
        ]
        for col, valor in enumerate(valores_calc, 1):
            celda = ws_calc.cell(idx, col, valor)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=col in {10, 11, 12})
            if col in {4} and isinstance(valor, date):
                celda.number_format = "DD/MM/YYYY"
            if col == 7:
                celda.number_format = '"$"#,##0'
            if col in {4, 5, 6, 7, 8, 9}:
                celda.fill = gris

    widths_calc = {
        "A": 8,
        "B": 18,
        "C": 16,
        "D": 16,
        "E": 14,
        "F": 12,
        "G": 24,
        "H": 16,
        "I": 14,
        "J": 24,
        "K": 24,
        "L": 22,
    }
    for col, width in widths_calc.items():
        ws_calc.column_dimensions[col].width = width
    ws_calc.row_dimensions[1].height = 28
    ws_calc.row_dimensions[4].height = 38
    ws_calc.freeze_panes = "A5"

    widths = {
        "A": 12,
        "B": 18,
        "C": 13,
        "D": 15,
        "E": 18,
        "F": 14,
        "G": 15,
        "H": 28,
        "I": 20,
        "J": 55,
        "K": 34,
        "L": 18,
        "M": 18,
        "N": 24,
        "O": 18,
        "P": 24,
        "Q": 20,
        "R": 15,
        "S": 24,
        "T": 14,
        "V": 12,
        "W": 26,
        "X": 15,
        "Y": 14,
        "Z": 18,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[7].height = 48
    ws.freeze_panes = "A8"

    for consecutivo, fila in enumerate(filas_ordenadas, 1):
        _agregar_ficha_contrato(wb, fila, consecutivo, profesional)

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _set_paragraph_text(paragraph, texto: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = texto
    else:
        paragraph.add_run(texto)


def _copiar_formato_run(origen, destino) -> None:
    rpr = origen._r.rPr if origen is not None else None
    if rpr is not None:
        destino._r.insert(0, copy.deepcopy(rpr))


def _set_paragraph_runs(paragraph, partes: list[tuple[str, bool | None, int | None]]) -> None:
    referencias = list(paragraph.runs)
    for run in referencias:
        paragraph._p.remove(run._r)
    for texto, negrilla, indice_referencia in partes:
        run = paragraph.add_run(texto)
        if indice_referencia is not None and indice_referencia < len(referencias):
            _copiar_formato_run(referencias[indice_referencia], run)
        if negrilla is not None:
            run.bold = negrilla


def _quitar_subrayado_xml(elemento) -> None:
    from docx.oxml.ns import qn

    for rpr in elemento.findall(".//" + qn("w:rPr")):
        for underline in list(rpr.findall(qn("w:u"))):
            rpr.remove(underline)


def _ajustar_subrayado_rpr(rpr, subrayado: bool) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for underline in list(rpr.findall(qn("w:u"))):
        rpr.remove(underline)
    if subrayado:
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(underline)


def _copiar_parrafo_xml(
    sample_p,
    texto: str,
    *,
    sin_subrayado: bool = False,
    con_subrayado: bool = False,
):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = copy.deepcopy(sample_p)
    if sin_subrayado:
        _quitar_subrayado_xml(p)
    primer_rpr = None
    for r in p.findall(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            primer_rpr = copy.deepcopy(rpr)
            break
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    r = OxmlElement("w:r")
    if primer_rpr is None and (sin_subrayado or con_subrayado):
        primer_rpr = OxmlElement("w:rPr")
    if primer_rpr is not None:
        if sin_subrayado:
            _ajustar_subrayado_rpr(primer_rpr, False)
        elif con_subrayado:
            _ajustar_subrayado_rpr(primer_rpr, True)
        r.append(primer_rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = texto
    r.append(t)
    p.append(r)
    return p


def _insertar_parrafo(
    anchor,
    sample_p,
    texto: str,
    *,
    sin_subrayado: bool = False,
    con_subrayado: bool = False,
) -> None:
    anchor.addprevious(
        _copiar_parrafo_xml(
            sample_p,
            texto,
            sin_subrayado=sin_subrayado,
            con_subrayado=con_subrayado,
        )
    )


def _insertar_blanco(anchor, sample_p) -> None:
    anchor.addprevious(copy.deepcopy(sample_p))


def _aplicar_fuente_run(run) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = "Garamond"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Garamond")
    rfonts.set(qn("w:hAnsi"), "Garamond")
    rfonts.set(qn("w:eastAsia"), "Garamond")
    run.font.size = Pt(11)
    run.font.highlight_color = None


def _set_cell_text(cell, texto: str) -> None:
    cell.text = str(texto or "")
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            _aplicar_fuente_run(run)


def _aplicar_fuente_parrafo(paragraph, *, negrilla: bool | None = None, subrayado: bool | None = None) -> None:
    for run in paragraph.runs:
        _aplicar_fuente_run(run)
        if negrilla is not None:
            run.bold = negrilla
        if subrayado is not None:
            run.underline = subrayado


def _aplicar_formato_tabla_solicitud(table) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    table.autofit = False
    widths = [Inches(1.45), Inches(4.25), Inches(0.42), Inches(0.42)]
    for row_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Inches(0.29)
        if row_idx == 4:
            row.height = Inches(1.18)
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if col_idx < len(widths):
                cell.width = widths[col_idx]
            for paragraph in cell.paragraphs:
                if row_idx in {0, 1} or col_idx in {1, 2, 3}:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _aplicar_fuente_parrafo(paragraph)


def _marcar_se_ajusta(table, row_idx: int, se_ajusta: bool | None) -> None:
    _set_cell_text(table.rows[row_idx].cells[2], "")
    _set_cell_text(table.rows[row_idx].cells[3], "")
    if se_ajusta is None:
        return
    _set_cell_text(table.rows[row_idx].cells[2 if se_ajusta else 3], "X")


def _observaciones_validacion(fila: dict, validacion: dict) -> list[str]:
    observaciones: list[str] = []
    fecha_fin_solicitud = parsear_fecha(fila.get("fecha_terminacion_inicial"))
    fecha_final_solicitud = parsear_fecha(fila.get("fecha_terminacion_final"))
    valor_adicion = _a_numero(fila.get("valor_adicion_texto")) or _a_numero(fila.get("valor_adicion"))

    fecha_fin_calculada = validacion.get("fecha_fin_calculada")
    if fecha_fin_solicitud and fecha_fin_calculada and fecha_fin_solicitud != fecha_fin_calculada:
        observaciones.append(f"Fecha de terminación: Sería el {formato_fecha_larga(fecha_fin_calculada)}.")

    prorroga_texto = _texto(fila.get("prorroga_solicitada")).strip()
    prorroga_ajustada = _texto(validacion.get("prorroga_ajustada_texto")).strip()
    if prorroga_texto and prorroga_ajustada and _normalizar(prorroga_texto) != _normalizar(prorroga_ajustada):
        observaciones.append(f"Prorroga: Sería por {prorroga_ajustada}.")

    fecha_final_ajustada = validacion.get("fecha_final_ajustada")
    if fecha_final_solicitud and fecha_final_ajustada and fecha_final_solicitud != fecha_final_ajustada:
        observaciones.append(
            f"Fecha de terminación con la prórroga: Sería el {formato_fecha_larga(fecha_final_ajustada)}."
        )

    valor_adicion_teorico = validacion.get("valor_adicion_teorico")
    if (
        valor_adicion is not None
        and valor_adicion_teorico is not None
        and abs(valor_adicion - valor_adicion_teorico) > 1
    ):
        observaciones.append(f"Adición: Sería por {formato_moneda(valor_adicion_teorico)}.")
    return observaciones


def _llenar_tabla_solicitud(table, fila: dict) -> None:
    validacion = _calcular_validacion_solicitud(fila)
    valores = {
        2: _texto(fila.get("contrato")),
        3: _texto(fila.get("contratista")),
        4: _texto(fila.get("objeto")),
        5: _texto(fila.get("fecha_inicio")),
        6: PLAZO_SOLICITUD_WORD,
        7: _texto(fila.get("valor_inicial_texto")) or formato_moneda(fila.get("valor_inicial")),
        8: _texto(fila.get("fecha_terminacion_inicial")),
        9: _texto(fila.get("prorroga_solicitada")),
        10: _texto(fila.get("valor_adicion_texto")) or formato_moneda(fila.get("valor_adicion")),
        11: _texto(fila.get("fecha_terminacion_final")),
    }
    for row_idx, valor in valores.items():
        _set_cell_text(table.rows[row_idx].cells[1], valor)
    _marcar_se_ajusta(
        table,
        8,
        validacion["se_ajusta_fecha_fin"]
        if validacion.get("fecha_fin_calculada") and parsear_fecha(fila.get("fecha_terminacion_inicial"))
        else None,
    )
    _marcar_se_ajusta(
        table,
        9,
        validacion["se_ajusta_prorroga"] if _texto(fila.get("prorroga_solicitada")).strip() else None,
    )
    _marcar_se_ajusta(
        table,
        10,
        validacion["se_ajusta_adicion"]
        if validacion.get("valor_adicion_teorico") is not None
        and (_a_numero(fila.get("valor_adicion_texto")) is not None or _a_numero(fila.get("valor_adicion")) is not None)
        else None,
    )
    _marcar_se_ajusta(
        table,
        11,
        validacion["se_ajusta_fecha_final"]
        if validacion.get("fecha_final_calculada") and parsear_fecha(fila.get("fecha_terminacion_final"))
        else None,
    )
    _aplicar_formato_tabla_solicitud(table)


def _quitar_resaltados(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _aplicar_fuente_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _aplicar_fuente_run(run)


def _cargo_localidad(localidad: str, supervisor: str) -> str:
    nombre = _normalizar(supervisor)
    cargo = "Alcaldesa" if any(p in nombre.split() for p in ("maria", "angela", "angelica", "diana", "paola", "catherine", "claudia", "andrea", "luisa")) else "Alcalde(sa)"
    return f"{cargo} Local de {localidad}"


def _texto_observacion_vigencia(fila: dict, validacion: dict | None = None) -> str:
    fin_final = None
    if validacion:
        fin_final = validacion.get("fecha_final_ajustada")
    fin_final = fin_final or parsear_fecha(fila.get("fecha_terminacion_final"))
    if not fin_final or fin_final <= CIERRE_VIGENCIA_FISCAL_2026:
        return ""
    return (
        "Prorroga: Supera la vigencia fiscal 2026, si bien la regla general impide la ejecución "
        "contractual más allá de la vigencia fiscal en curso, existe habilitación legal para su prórroga "
        "excepcional, siempre y cuando se sustente en necesidades claramente justificadas, se respeten "
        "los límites establecidos para la contratación estatal y se observen estrictamente los procedimientos "
        "presupuestales correspondientes."
    )


def generar_documento_localidad(
    filas: list[dict],
    plantilla_docx: Path,
    fecha_respuesta: date,
    profesional: str,
) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table

    filas = ordenar_filas(filas)
    localidad = _canon_localidad(filas[0].get("localidad", ""))
    supervisor = _texto(filas[0].get("supervisor")).upper()
    sipses = []
    for fila in filas:
        sipse = _texto(fila.get("sipse")).strip()
        if sipse and sipse not in sipses:
            sipses.append(sipse)
    sipses_txt = ", ".join(sipses) if sipses else "SIN NÚMERO SIPSE"
    radicado = _texto(filas[0].get("radicado_salida")).strip()
    fecha_salida = parsear_fecha(filas[0].get("fecha_salida"))

    doc = Document(str(plantilla_docx))
    paragraphs = doc.paragraphs
    first_solicitud = next((p for p in paragraphs if p.text.strip().upper().startswith("SOLICITUD")), None)
    recomendaciones = next((p for p in paragraphs if "RECOMENDACIONES GENERALES" in p.text.upper()), None)
    if first_solicitud is None or recomendaciones is None or not doc.tables:
        raise ValueError("La plantilla Word no tiene la estructura esperada.")

    sample_heading = copy.deepcopy(first_solicitud._p)
    sample_blank = copy.deepcopy(doc.paragraphs[paragraphs.index(first_solicitud) + 1]._p)
    sample_normal = copy.deepcopy(next((p._p for p in paragraphs if p.text.strip().lower().startswith("prorroga:")), first_solicitud._p))
    sample_intro = copy.deepcopy(next((p._p for p in paragraphs if p.text.strip().lower().startswith("en cuanto")), first_solicitud._p))
    sample_table = copy.deepcopy(doc.tables[0]._tbl)

    for p in doc.paragraphs:
        texto = p.text.strip()
        if texto.startswith("Bogotá D.C."):
            _set_paragraph_text(p, f"Bogotá D.C., {formato_fecha_larga(fecha_respuesta)}")
        elif texto.startswith("PARA:"):
            _set_paragraph_runs(
                p,
                [
                    ("PARA:", True, 0),
                    ("\t", False, 2),
                    (supervisor, True, 3),
                ],
            )
        elif texto.startswith("Alcalde") or texto.startswith("Alcaldesa"):
            _set_paragraph_runs(
                p,
                [
                    ("\t", False, 0),
                    (_cargo_localidad(localidad, supervisor), False, 1),
                ],
            )
        elif texto.startswith("ASUNTO:"):
            asunto = (
                "RESPUESTA A SOLICITUDES SIPSE: "
                f"{sipses_txt}, SOBRE ADICIONES Y PRÓRROGAS A CONTRATOS DE PRESTACIÓN "
                "DE SERVICIOS PROFESIONALES Y DE APOYO A LA GESTIÓN."
            )
            _set_paragraph_runs(
                p,
                [
                    ("ASUNTO:", True, 0),
                    ("\t", False, 1),
                    (asunto, False, 2),
                ],
            )
        elif "mediante memorando de radicado" in texto:
            fecha_txt = formato_fecha_larga(fecha_salida) or formato_fecha_larga(fecha_respuesta)
            if radicado:
                _set_paragraph_text(
                    p,
                    (
                        "Importante resaltar que la Dirección para la Gestión del Desarrollo Local "
                        f"a través del radicado {radicado} del {fecha_txt}, emitió "
                        "“RECOMENDACIONES, SUGERENCIAS Y ALERTAS TEMPRANAS” con la finalidad de dar "
                        "seguridad técnica y jurídica a los Fondos de Desarrollo Local."
                    ),
                )
        elif texto.startswith("Proyectó:"):
            _set_paragraph_text(p, f"Proyectó: {profesional} - Profesional DGDL")

    body = doc.element.body
    borrar = False
    for child in list(body):
        if child is first_solicitud._p:
            borrar = True
        if child is recomendaciones._p:
            borrar = False
        if borrar:
            body.remove(child)

    anchor = recomendaciones._p
    for fila in filas:
        validacion = _calcular_validacion_solicitud(fila)
        _insertar_parrafo(anchor, sample_heading, f"SOLICITUD {_texto(fila.get('sipse')) or _texto(fila.get('archivo'))}".strip())
        _insertar_blanco(anchor, sample_blank)
        tbl_xml = copy.deepcopy(sample_table)
        tabla = Table(tbl_xml, doc)
        _llenar_tabla_solicitud(tabla, fila)
        anchor.addprevious(tbl_xml)
        _insertar_blanco(anchor, sample_blank)
        _insertar_parrafo(anchor, sample_normal, "Observaciones respecto de se ajusta Si/No:", con_subrayado=True)
        _insertar_blanco(anchor, sample_blank)
        observaciones = _observaciones_validacion(fila, validacion)
        observacion_vigencia = _texto_observacion_vigencia(fila, validacion)
        if observacion_vigencia:
            observaciones.append(observacion_vigencia)
        for observacion in observaciones:
            _insertar_parrafo(anchor, sample_normal, f"•\t{observacion}", sin_subrayado=True)
            _insertar_blanco(anchor, sample_blank)
        _insertar_parrafo(
            anchor,
            sample_intro,
            "En cuanto a la información que registra a la fecha en el contrato electrónico se observa que:",
            con_subrayado=True,
        )
        _insertar_blanco(anchor, sample_blank)
        _insertar_parrafo(
            anchor,
            sample_normal,
            "a.\tEl contratista ha cargado la documentación e información que evidencia su ejecución contractual.",
            sin_subrayado=True,
        )
        _insertar_blanco(anchor, sample_blank)
        _insertar_blanco(anchor, sample_blank)

    _quitar_resaltados(doc)
    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue()


def generar_zip_asistencia(
    filas: Iterable[dict],
    plantilla_docx: Path,
    fecha_respuesta: date,
    profesional: str,
) -> tuple[bytes, str, list[dict]]:
    filas_ordenadas = ordenar_filas(filas)
    if not filas_ordenadas:
        raise ValueError("No hay solicitudes para generar.")

    resumen: list[dict] = []
    salida = BytesIO()
    with zipfile.ZipFile(salida, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        localidades = []
        for fila in filas_ordenadas:
            localidad = _canon_localidad(fila.get("localidad", ""))
            if localidad and localidad not in localidades:
                localidades.append(localidad)
        for localidad in localidades:
            grupo = [f for f in filas_ordenadas if _canon_localidad(f.get("localidad", "")) == localidad]
            docx = generar_documento_localidad(grupo, plantilla_docx, fecha_respuesta, profesional)
            nombre = f"Respuesta asistencia tecnica {localidad} {fecha_respuesta.year}.docx"
            zf.writestr(nombre, docx)
            resumen.append({"localidad": localidad, "solicitudes": len(grupo), "archivo": nombre})

        excel = generar_excel_asistencia(filas_ordenadas, profesional=profesional)
        zf.writestr("Matriz asistencia tecnica CPS.xlsx", excel)

    localidades_txt = ",".join(_canon_localidad(f.get("localidad", "")) for f in filas_ordenadas)
    nombre_zip = f"Asistencia tecnica CPS {fecha_respuesta.year} ({_nombre_archivo_seguro(localidades_txt)[:80]}).zip"
    return salida.getvalue(), nombre_zip, resumen
