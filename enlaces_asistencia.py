"""Cruce de enlaces del archivo CPS, sin consultar ni modificar datos de la solicitud."""
from io import BytesIO
import re
from urllib.parse import urlsplit

from asistencia_tecnica import (
    LOCALIDADES_RADICADO, _canon_localidad, _normalizar,
    _numero_contrato_corto, _texto, ordenar_filas,
)


NOMBRE_EXCEL_ENLACES = "Enlaces a contratos CPS.xlsx"


def _localidad_entidad(entidad):
    texto = _normalizar(entidad)
    for localidad in sorted(LOCALIDADES_RADICADO, key=len, reverse=True):
        nombre = re.sub(r"^(la|los) ", "", _normalizar(localidad))
        if re.search(r"(?<!\w)" + re.escape(nombre) + r"(?!\w)", texto):
            return localidad
    return ""


def _contrato_completo(numero):
    return re.sub(r"[^a-z0-9]", "", _normalizar(numero))


def _url_valida(valor):
    texto = _texto(valor).strip()
    try:
        url = urlsplit(texto)
        return texto if url.scheme in {"https", "http"} and url.hostname and not re.search(r"\s", texto) else ""
    except ValueError:
        return ""


def leer_base_enlaces(contenido: bytes) -> dict:
    """Indexa únicamente entidad, número, identificador y URL; no importa otros datos."""
    from openpyxl import load_workbook

    indice = {}
    encontrada = False
    wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    try:
        for hoja in wb:
            filas = hoja.iter_rows(values_only=True)
            encabezados = {_normalizar(v): i for i, v in enumerate(next(filas, ()))}
            requeridos = ("nombre entidad", "numero del contrato", "urlproceso")
            if not all(c in encabezados for c in requeridos):
                continue
            encontrada = True
            for fila in filas:
                def valor(campo):
                    pos = encabezados.get(campo)
                    return _texto(fila[pos]).strip() if pos is not None and pos < len(fila) else ""
                localidad = _localidad_entidad(valor("nombre entidad"))
                numero = valor("numero del contrato")
                corto = _numero_contrato_corto(numero)
                if not localidad or not numero:
                    continue
                registro = (numero, valor("id contrato"), _url_valida(valor("urlproceso")), valor("nombre entidad"))
                grupo = indice.setdefault((localidad, corto), [])
                if registro not in grupo:
                    grupo.append(registro)
    finally:
        wb.close()
    if not encontrada:
        raise ValueError("El Excel debe incluir Nombre Entidad, Número del Contrato y URLProceso.")
    return indice


def cruzar_enlaces(filas, indice: dict | None):
    resultado = []
    for original in filas:
        fila = dict(original)
        fila.update(link_secop="", entidad_enlace="", estado_enlace="Sin base de enlaces")
        if indice is not None:
            numero = _texto(fila.get("contrato"))
            clave = (_canon_localidad(fila.get("localidad", "")), _numero_contrato_corto(numero))
            candidatos = indice.get(clave, [])
            candidatos = [c for c in candidatos if _contrato_completo(c[0]) == _contrato_completo(numero)]
            # Una base duplicada no debe crear ambigüedad, pero contratos distintos sí.
            identidades = {(c[1] or _contrato_completo(c[0]), c[2]) for c in candidatos}
            if not identidades:
                fila["estado_enlace"] = "No encontrado en la base"
            elif len(identidades) > 1:
                fila["estado_enlace"] = "Revisar: varias coincidencias"
            else:
                _, enlace = next(iter(identidades))
                fila["link_secop"] = enlace
                fila["entidad_enlace"] = candidatos[0][3]
                fila["estado_enlace"] = "Encontrado" if enlace else "Sin URL válida en la base"
        resultado.append(fila)
    return resultado


def generar_excel_enlaces(filas) -> bytes:
    # Se integra con el mismo motor Excel de la aplicación desplegada.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Enlaces CPS"
    ws.sheet_view.showGridLines = False
    ws.append(["Localidad del Word", "No. en Word", "SIPSE", "Contrato (columna L)", "Entidad (columna A)", "Enlace SECOP (columna AW)", "Resultado del cruce"])
    anterior = None
    numero = 0
    for fila in ordenar_filas(filas):
        localidad = _canon_localidad(fila.get("localidad", ""))
        if not localidad:
            continue  # El generador Word tampoco crea documento sin localidad.
        numero = numero + 1 if localidad == anterior else 1
        anterior = localidad
        enlace = _url_valida(fila.get("link_secop"))
        valores = [localidad, numero, _texto(fila.get("sipse")), _texto(fila.get("contrato")),
                   _texto(fila.get("entidad_enlace")), enlace,
                   _texto(fila.get("estado_enlace")) or ("Encontrado" if enlace else "Sin enlace")]
        ws.append(valores)
        for celda in ws[ws.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"
            celda.font = Font(name="Garamond", size=11)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            if numero % 2 == 0:
                celda.fill = PatternFill("solid", fgColor="F2F2F2")
        for columna in (2, 3):
            ws.cell(ws.max_row, columna).alignment = Alignment(horizontal="center", vertical="center")
        if enlace:
            celda = ws.cell(ws.max_row, 6)
            celda.hyperlink = enlace
            celda.font = Font(name="Garamond", size=11, color="0563C1", underline="single")
        ws.row_dimensions[ws.max_row].height = 66
    for celda in ws[1]:
        celda.font = Font(name="Garamond", size=11, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="C00000")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for columna, ancho in zip("ABCDEFG", [22, 14, 14, 30, 38, 75, 32]):
        ws.column_dimensions[columna].width = ancho
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
